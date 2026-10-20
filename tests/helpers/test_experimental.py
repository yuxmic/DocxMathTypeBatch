"""生成不含私人资料的区域混合样本，并独立检查转换内容与容器属性。"""
import json
import posixpath
import re
import sys
import unicodedata
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
V = "urn:schemas-microsoft-com:vml"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "m": M}


def math_run(value, display=False):
    equation = parse_xml(f'<m:oMath xmlns:m="{M}"><m:f><m:num><m:r><m:t>x+{value}</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>')
    if not display:
        return equation
    wrapper = OxmlElement("m:oMathPara")
    wrapper.append(equation)
    return wrapper


def fill(paragraph, label, value):
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.add_run(f"{label} $x+{value}$ / ")
    paragraph._p.append(math_run(value + 100))
    paragraph.add_run(" tail")


def fixture(folder):
    doc = Document()
    doc.core_properties.author = "Synthetic test fixture"
    doc.core_properties.last_modified_by = ""
    doc.add_paragraph("Basic inline $a+1$ tail")
    fill(doc.add_paragraph(), "Body", 11)
    table = doc.add_table(rows=2, cols=2)
    for index, cell in enumerate(table._cells):
        fill(cell.paragraphs[0], f"Cell{index}", 21 + index)
    fill(doc.sections[0].header.paragraphs[0], "Header", 31)
    fill(doc.sections[0].footer.paragraphs[0], "Footer", 41)
    # 文本框使用 Word 可识别的 VML 容器，与实际论文无关。
    pict = OxmlElement("w:pict")
    shape = etree.SubElement(pict, f"{{{V}}}shape", id="SyntheticTextBox", style="width:240pt;height:90pt", type="#_x0000_t202")
    textbox = etree.SubElement(shape, f"{{{V}}}textbox")
    content = etree.SubElement(textbox, qn("w:txbxContent"))
    p = etree.SubElement(content, qn("w:p"))
    run = etree.SubElement(p, qn("w:r"))
    etree.SubElement(run, qn("w:t")).text = "Textbox $x+51$ / "
    p.append(math_run(151))
    etree.SubElement(etree.SubElement(p, qn("w:r")), qn("w:t")).text = " tail"
    doc.add_paragraph().add_run()._r.append(pict)
    for name in ("footnote", "endnote"):
        reference = OxmlElement(f"w:{name}Reference")
        reference.set(qn("w:id"), "1")
        doc.add_paragraph(f"{name} reference ").add_run()._r.append(reference)
    target = folder / "input.docx"
    doc.save(target)
    with ZipFile(target) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    relations = etree.fromstring(parts["word/_rels/document.xml.rels"])
    types = etree.fromstring(parts["[Content_Types].xml"])
    for name, value in (("footnote", 61), ("endnote", 71)):
        root = etree.Element(qn(f"w:{name}s"), nsmap={"w": W, "m": M})
        for ident, kind, tag in ((-1, "separator", "separator"), (0, "continuationSeparator", "continuationSeparator")):
            note = etree.SubElement(root, qn(f"w:{name}"), {qn("w:id"): str(ident), qn("w:type"): kind})
            etree.SubElement(etree.SubElement(etree.SubElement(note, qn("w:p")), qn("w:r")), qn(f"w:{tag}"))
        note = etree.SubElement(root, qn(f"w:{name}"), {qn("w:id"): "1"})
        p = etree.SubElement(note, qn("w:p"))
        etree.SubElement(etree.SubElement(p, qn("w:r")), qn(f"w:{name}Ref"))
        etree.SubElement(etree.SubElement(p, qn("w:r")), qn("w:t")).text = f" {name} $x+{value}$ / "
        p.append(math_run(value + 100))
        etree.SubElement(etree.SubElement(p, qn("w:r")), qn("w:t")).text = " tail"
        parts[f"word/{name}s.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        etree.SubElement(relations, f"{{{REL}}}Relationship", Id=f"rIdSynthetic{name}", Type=f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{name}s", Target=f"{name}s.xml")
        etree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", PartName=f"/word/{name}s.xml", ContentType=f"application/vnd.openxmlformats-officedocument.wordprocessingml.{name}s+xml")
    parts["word/_rels/document.xml.rels"] = etree.tostring(relations, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
    for name in list(parts):
        if name.startswith("word/") and name.endswith(".xml"):
            root = etree.fromstring(parts[name])
            for node in root.xpath("//w:t", namespaces=NS):
                if node.text and node.text != node.text.strip():
                    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    # 额外准备无基本目标、非 ASCII 保留和行内/行外筛选样本。
    for variant in ("experimental-only", "nonascii-input", "rollback-input", "scope-input"):
        changed = dict(parts)
        root = etree.fromstring(changed["word/document.xml"])
        if variant == "experimental-only":
            for node in root.xpath("/w:document/w:body/w:p/w:r/w:t", namespaces=NS):
                if node.text:
                    node.text = node.text.replace("$a+1$", "kept").replace("$x+11$", "kept")
            # 隔离样本只保留额外位置的公式，正文 OMML 现在属于基本功能。
            for equation in root.xpath("/w:document/w:body/w:p//m:oMath[not(ancestor::w:txbxContent)]", namespaces=NS):
                equation.getparent().remove(equation)
        elif variant in ("nonascii-input", "rollback-input"):
            for node in root.xpath("//w:t", namespaces=NS):
                if node.text:
                    node.text = node.text.replace("$x+21$", r"$\text{测试}$" if variant == "nonascii-input" else "$experimental_rollback$")
        else:
            body = root.find(qn("w:body"))
            p = etree.Element(qn("w:p"))
            p.append(math_run(909, True))
            body.insert(2, p)
            cell = root.xpath("/w:document/w:body/w:tbl/w:tr/w:tc", namespaces=NS)[0]
            p = etree.SubElement(cell, qn("w:p"))
            etree.SubElement(etree.SubElement(p, qn("w:r")), qn("w:t")).text = r"\[z+88\]"
            content = root.xpath("//w:txbxContent", namespaces=NS)[0]
            etree.SubElement(content, qn("w:p")).append(math_run(908, True))
        changed["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        with ZipFile(folder / f"{variant}.docx", "w", ZIP_DEFLATED) as archive:
            for name, data in changed.items():
                archive.writestr(name, data)
    print("EXPERIMENTAL_FIXTURE_CREATED")


def no_targets_fixture(folder):
    """正文保留公式，已启用额外区域只留普通文字，用于验证存在性预检短路。"""
    source = folder / "input.docx"
    if not source.exists():
        fixture(folder)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    # 表格和文本框位于正文部件内，只清理这两类子容器，正文段落公式保留。
    root = etree.fromstring(parts["word/document.xml"])
    for container in root.xpath("//w:tbl", namespaces=NS) + root.xpath("//w:txbxContent", namespaces=NS):
        for node in container.xpath(".//w:t", namespaces=NS):
            if node.text:
                node.text = TEX.sub("", node.text)
        for equation in container.xpath(".//m:oMath", namespaces=NS):
            equation.getparent().remove(equation)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    # 页眉页脚与脚注尾注是独立部件，整部件删除 LaTeX 片段和 OMML。
    for name in list(parts):
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        if not (name.startswith(("word/header", "word/footer")) or name.endswith(("footnotes.xml", "endnotes.xml"))):
            continue
        root = etree.fromstring(parts[name])
        for node in root.xpath("//w:t", namespaces=NS):
            if node.text:
                node.text = TEX.sub("", node.text)
        for equation in root.xpath("//m:oMath", namespaces=NS):
            equation.getparent().remove(equation)
        parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    target = folder / "no-targets-input.docx"
    with ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    # 生成后自检：只剩正文三个目标，启用区域既没有定界符也没有 OMML。
    remaining = target_count(target)
    assert remaining == 3, remaining
    print("EXPERIMENTAL_NO_TARGETS_FIXTURE_CREATED")


def no_changes_fixture(folder):
    """额外区域全部换成非 ASCII 公式并删除 OMML：有候选，但实验转换必然全部失败。"""
    source = folder / "input.docx"
    if not source.exists():
        fixture(folder)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    def replace(container):
        # 保留 $...$ 定界符作为候选，内容换成导入前就会失败的非 ASCII 公式。
        # 用函数替换避免 re.sub 把替换串里的 \t 当成制表符。
        for node in container.xpath(".//w:t", namespaces=NS):
            if node.text:
                node.text = re.sub(r"\$x\+\d+\$", lambda _: r"$\text{测试}$", node.text)
        for equation in container.xpath(".//m:oMath", namespaces=NS):
            equation.getparent().remove(equation)
    # 表格和文本框位于正文部件内；正文段落公式保持不变。
    root = etree.fromstring(parts["word/document.xml"])
    for container in root.xpath("//w:tbl", namespaces=NS) + root.xpath("//w:txbxContent", namespaces=NS):
        replace(container)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    # 页眉页脚与脚注尾注是独立部件，整部件替换并删除 OMML。
    for name in list(parts):
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        if not (name.startswith(("word/header", "word/footer")) or name.endswith(("footnotes.xml", "endnotes.xml"))):
            continue
        root = etree.fromstring(parts[name])
        replace(root)
        parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    target = folder / "no-changes-input.docx"
    with ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    # 生成后自检：正文 3 个可转换目标保留，额外区域 9 个非 ASCII 候选计数在内。
    remaining = target_count(target)
    assert remaining == 12, remaining
    print("EXPERIMENTAL_NO_CHANGES_FIXTURE_CREATED")


def document_parts(path):
    result = {}
    with ZipFile(path) as archive:
        document = etree.fromstring(archive.read("word/document.xml"))
        relationships = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        targets = {node.get("Id"): posixpath.normpath("word/" + node.get("Target")) for node in relationships}
        owners = {}
        for section, properties in enumerate(document.xpath("//w:sectPr[not(ancestor::w:sectPrChange) and not(ancestor::w:txbxContent)]", namespaces=NS), 1):
            for reference in properties:
                kind = etree.QName(reference).localname
                if kind not in ("headerReference", "footerReference"):
                    continue
                relation = reference.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                owners.setdefault(targets[relation], f"{kind[:-9]}:{section}:{reference.get(qn('w:type'))}")
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            root = etree.fromstring(archive.read(name))
            # 当前 Word 生成的 DrawingML/VML 兼容副本只选一个分支，不能重复计算公式。
            for alternate in root.xpath("//*[local-name()='AlternateContent' and namespace-uri()='http://schemas.openxmlformats.org/markup-compatibility/2006']")[::-1]:
                choices = alternate.xpath("*[local-name()='Choice']")
                chosen = choices[0] if choices else alternate.xpath("*[local-name()='Fallback']")[0]
                parent = alternate.getparent()
                index = parent.index(alternate)
                for child in list(chosen):
                    parent.insert(index, child)
                    index += 1
                parent.remove(alternate)
            if not root.xpath("//w:p", namespaces=NS):
                continue
            if name.startswith(("word/header", "word/footer")):
                if name not in owners or not root.xpath("//w:t[text()] | //m:oMath | //w:instrText | //w:drawing | //w:pict | //*[local-name()='OLEObject']", namespaces=NS):
                    continue
                name = owners[name]
            result[name] = root
    return result


def inventory(path):
    return {name: {"OMML": len(root.xpath("//m:oMath", namespaces=NS)), "MathType": len(root.xpath("//*[local-name()='OLEObject' and @ProgID='Equation.DSMT4']")), "Text": "".join(root.xpath("//w:t/text()", namespaces=NS))} for name, root in document_parts(path).items()}


TEX = re.compile(r"\$[^$]+\$|\\\[[\s\S]*?\\\]|\\\([\s\S]*?\\\)")


def target_count(path):
    return sum(part["OMML"] + len(TEX.findall(part["Text"])) for part in inventory(path).values())


def check(source, output, case, baseline):
    before, after = inventory(baseline), inventory(output)
    counts = {"basic": 3, "tables": 11, "textboxes": 5, "headers": 7, "notes": 7, "all": 21, "complex-basic": 3, "noformat-basic": 0, "noformat-default": 18, "noformat-advanced": 18, "nonascii": 10, "rollback": 10, "scope-basic": 4, "scope-inline": 21, "scope-display": 3, "complex-notes": 8}
    expected = target_count(source) if case == "complex-all" else counts[case]
    assert sum(part["MathType"] for part in after.values()) == expected, after
    remaining = {"basic": 9, "complex-basic": 9, "tables": 5, "textboxes": 8, "headers": 7, "notes": 7, "nonascii": 5, "rollback": 5, "complex-notes": 7, "scope-basic": 10, "all": 0, "complex-all": 0, "noformat-default": 0, "noformat-advanced": 0, "scope-inline": 2, "scope-display": 10}.get(case, sum(part["OMML"] for part in inventory(source).values()))
    assert sum(part["OMML"] for part in after.values()) == remaining, after
    if case == "nonascii":
        assert r"\text{测试}" in after["word/document.xml"]["Text"]
    if case == "rollback":
        assert "$experimental_rollback$" in after["word/document.xml"]["Text"]
    for name in before:
        if not name.startswith(("header:", "footer:")) and name not in ("word/footnotes.xml", "word/endnotes.xml"):
            continue
        enabled = case in {"all", "complex-all", "noformat-default", "noformat-advanced", "scope-inline", "scope-display"} or (case == "headers" and name.startswith(("header:", "footer:"))) or (case in {"notes", "complex-notes"} and "notes.xml" in name)
        if not enabled:
            assert before[name] == after[name], (name, before[name], after[name])
    # 可见普通文字、表格尺寸及段落属性必须保留；不以对象数量代替排版检查。
    before_parts, after_parts = document_parts(baseline), document_parts(output)
    for name in before:
            assert name in after_parts, ("logical region lost", name)
            original, saved = before_parts[name], after_parts[name]
            for tag in ("tblPr", "trPr", "tcPr"):
                def props(root):
                    return [etree.tostring(node, method="c14n", exclusive=True) for node in root.xpath(f"//w:{tag}", namespaces=NS)]
                assert props(original) == props(saved), (case, name, tag)
            for label in ("Basic inline", "Body", "Cell0", "Cell1", "Cell2", "Cell3", "Header", "Footer", "Textbox", "footnote", "endnote"):
                original_text = "".join(original.xpath("//w:t/text()", namespaces=NS))
                if label in original_text:
                    assert label in after[name]["Text"], (case, name, label)
    print(json.dumps({"Case": case, "Status": "PASSED", "MathType": expected, "RemainingOMML": remaining}))


def region_snapshot(path):
    """只统计表格、文本框、页眉页脚、脚注尾注区域，正文段落不参与比较。"""
    result = {}
    for name, root in document_parts(path).items():
        if name == "word/document.xml":
            groups = [("word/document.xml#table", root.xpath("//w:tbl", namespaces=NS))]
            boxes = root.xpath("//w:txbxContent", namespaces=NS)
            if boxes:
                groups.append(("word/document.xml#textbox", boxes))
        elif name.startswith(("header:", "footer:")) or name.endswith(("footnotes.xml", "endnotes.xml")):
            groups = [(name, [root])]
        else:
            continue
        for key, nodes in groups:
            result[key] = {"Text": "".join("".join(node.xpath(".//w:t/text()", namespaces=NS)) for node in nodes),
                           "OMML": sum(len(node.xpath(".//m:oMath", namespaces=NS)) for node in nodes),
                           "MathType": sum(len(node.xpath(".//*[local-name()='OLEObject' and @ProgID='Equation.DSMT4']")) for node in nodes)}
    return result


def check_no_targets(folder):
    """独立核对无目标短路：正文照常转换，启用区域保持原样，报告与日志一致。"""
    result = json.loads((folder / "no-targets.json").read_text(encoding="utf-8-sig"))
    assert result["Status"] == "COMPLETE" and result["Succeeded"] == 3 and not result["Issues"], result
    experimental = result["Experimental"]
    assert experimental["Enabled"] is True and experimental["Skipped"] == "NO_TARGETS", experimental
    assert experimental["Succeeded"] == experimental["FailedPreserved"] == experimental["RejectedPreserved"] == experimental["NotAttempted"] == 0, experimental
    assert experimental["Items"] == [] and experimental["Issues"] == [], experimental
    assert experimental["Options"] == {"Tables": True, "Textboxes": True, "HeadersFooters": True, "FootnotesEndnotes": True}, experimental
    before, after = region_snapshot(folder / "no-targets-input.docx"), region_snapshot(folder / "no-targets.docx")
    assert set(after) == set(before), (sorted(before), sorted(after))
    for name in before:
        assert after[name] == before[name], (name, before[name], after[name])
    # 正文三个目标确实转换，启用区域没有产生任何 MathType 对象。
    actual = inventory(folder / "no-targets.docx")
    assert sum(part["MathType"] for part in actual.values()) == 3, actual
    assert sum(part["OMML"] for part in actual.values()) == 0, actual
    log = (folder / "no-targets.log").read_text(encoding="utf-8")
    assert "UI_EXPERIMENTAL_SKIPPED NO_TARGETS" in log, log
    assert "EXPERIMENTAL_CONVERSION_VERIFIED" not in log, log
    assert "UI_EXPERIMENTAL_PHASE CONVERTING" not in log, log
    print("EXPERIMENTAL_NO_TARGETS_PASSED")


def check_no_changes(folder):
    """独立核对无改动短路：候选全部失败并保留原式，正文照常转换，不保存重开。"""
    result = json.loads((folder / "no-changes.json").read_text(encoding="utf-8-sig"))
    assert result["Status"] == "PARTIAL" and result["Succeeded"] == 3 and result["FailedPreserved"] > 0, result
    experimental = result["Experimental"]
    assert experimental["Enabled"] is True and experimental["Skipped"] == "NO_CHANGES", experimental
    assert experimental["Succeeded"] == 0 and len(experimental["Items"]) == result["FailedPreserved"] == experimental["FailedPreserved"], experimental
    assert all(item["Status"] == "FAILED_PRESERVED" for item in experimental["Items"]), experimental["Items"]
    assert experimental["Options"] == {"Tables": True, "Textboxes": True, "HeadersFooters": True, "FootnotesEndnotes": True}, experimental
    # 失败恢复后额外区域必须与候选原式完全一致，正文三个目标则照常转换。
    before, after = region_snapshot(folder / "no-changes-input.docx"), region_snapshot(folder / "no-changes.docx")
    assert set(after) == set(before), (sorted(before), sorted(after))
    for name in before:
        assert after[name] == before[name], (name, before[name], after[name])
    actual = inventory(folder / "no-changes.docx")
    assert sum(part["MathType"] for part in actual.values()) == 3, actual
    assert sum(part["OMML"] for part in actual.values()) == 0, actual
    log = (folder / "no-changes.log").read_text(encoding="utf-8")
    assert "EXPERIMENTAL_NO_CHANGES" in log and "UI_EXPERIMENTAL_SKIPPED NO_CHANGES" in log, log
    assert "EXPERIMENTAL_CONVERSION_VERIFIED" not in log, log
    assert "UI_EXPERIMENTAL_PHASE VERIFYING" not in log, log
    print("EXPERIMENTAL_NO_CHANGES_PASSED")


def check_reverse(source, output):
    after = inventory(output)
    assert sum(part["MathType"] for part in after.values()) == 0, after
    assert sum(part["OMML"] for part in after.values()) == 0, after
    for name, root in document_parts(source).items():
        text = "".join(root.xpath("//w:t/text()", namespaces=NS))
        expected = [value[1:-1] if value.startswith("$") else value[2:-2] for value in TEX.findall(text)]
        expected += [unicodedata.normalize("NFKC", "".join(math.xpath(".//m:num//m:t/text()", namespaces=NS))) for math in root.xpath("//m:oMath", namespaces=NS)]
        actual = re.sub(r"([_^])\{([^{}]+)\}", r"\1\2", after[name]["Text"].replace(" ", ""))
        for value in expected:
            assert len(re.findall(re.escape(value) + r"(?![0-9])", actual)) == expected.count(value), (name, value, actual)
    print("EXPERIMENTAL_REVERSE_CONTENT_VERIFIED")


def header_only_fixture(folder):
    """仅主页面页眉含公式，其他五个页眉页脚槽位不存在。"""
    doc = Document()
    doc.core_properties.author = "Synthetic test fixture"
    doc.core_properties.last_modified_by = ""
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_paragraph(r"\[a+1\]")
    doc.add_paragraph(r"\[b+2\]")
    doc.sections[0].header.paragraphs[0].text = "$h_1$"
    doc.save(folder / "header-only-input.docx")
    # 历史节属性不能被误认为第二个实际节，也不能生成虚拟页眉区域。
    from copy import deepcopy
    from docx.enum.section import WD_SECTION_START

    history = Document()
    history.core_properties.author = "Synthetic test fixture"
    history.core_properties.last_modified_by = ""
    history.add_paragraph(r"\[a+1\]")
    history.sections[0].header.paragraphs[0].text = "$h_1$"
    section = history.add_section(WD_SECTION_START.NEW_PAGE)
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = "$h_2$"
    history.add_paragraph(r"\[b+2\]")
    history.save(folder / "header-history-input.docx")
    with ZipFile(folder / "header-history-input.docx") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    section = root.xpath("/w:document/w:body/w:p/w:pPr/w:sectPr", namespaces=NS)[0]
    change = etree.SubElement(section, qn("w:sectPrChange"), {qn("w:id"): "1", qn("w:author"): "Synthetic test fixture", qn("w:date"): "2026-01-01T00:00:00Z"})
    previous = deepcopy(section)
    previous.remove(previous.find(qn("w:sectPrChange")))
    # 历史格式只保留旧节的排版属性；不复制页眉部件引用，否则 Word 会重复展开同一页眉。
    for reference in previous.xpath("w:headerReference | w:footerReference", namespaces=NS):
        previous.remove(reference)
    margins = previous.find(qn("w:pgMar"))
    margins.set(qn("w:left"), str(int(margins.get(qn("w:left"))) + 100))
    change.append(previous)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(folder / "header-history-input.docx", "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def check_header_only(folder, case):
    from native_number_text import native_visible_numbers

    result = json.loads((folder / f"{case}.json").read_text(encoding="utf-8-sig"))
    assert result["Status"] == "COMPLETE", result
    roots = document_parts(folder / f"{case}.docx")
    headers = [root for name, root in roots.items() if name.startswith("header:")]
    assert len(headers) == (2 if case == "header-history-first" else 1) and not any(name.startswith("footer:") for name in roots)
    numbers = native_visible_numbers(roots["word/document.xml"])
    start = 7 if case == "header-only-renumber" else 1
    # 第一节存在格式修订时，Word 将其正文作为修订范围，沿用既有保护规则保留。
    display_count = 1 if case == "header-history-first" else 2
    assert numbers == [f"({value})" for value in range(start, start + display_count)], numbers
    if case == "header-history-first":
        assert r"\[a+1\]" in "".join(roots["word/document.xml"].xpath("//w:t/text()", namespaces=NS))
    assert not native_visible_numbers(headers[0]), "header was numbered"
    if case == "header-only-reverse":
        assert result["Succeeded"] == 3
        header_text = "".join(headers[0].xpath("//w:t/text()", namespaces=NS)).replace(" ", "")
        assert "h_1" in re.sub(r"_\{([^{}]+)\}", r"_\1", header_text)
        assert not any(root.xpath("//*[local-name()='OLEObject']") for root in roots.values())
    else:
        assert all(len(header.xpath("//*[local-name()='OLEObject' and @ProgID='Equation.DSMT4']")) == 1 for header in headers)
        if case == "header-only-renumber":
            assert result["Targets"] == 0 and result["NumbersAdded"] == 0 and result["NumbersUpdated"] == 2
            # 页眉只剩已转换 MathType，没有 LaTeX 或 OMML 候选，实验阶段应短路。
            assert result["Experimental"]["Succeeded"] == 0 and result["Experimental"]["Skipped"] == "NO_TARGETS", result["Experimental"]
        else:
            assert result["Succeeded"] == 3 and result["NumbersAdded"] == display_count
    print(f"HEADER_ONLY_PASSED: {case}")


def review_fixture(folder):
    """覆盖未显示的已保存页眉，以及多张表格前后正文的最终报告位置。"""
    for name, slot, omml in (
        ('review-first-header', 'first_page_header', False),
        ('review-even-header', 'even_page_header', True),
        ('review-first-footer', 'first_page_footer', True),
        ('review-even-footer', 'even_page_footer', False),
    ):
        doc = Document()
        doc.core_properties.author = 'Synthetic test fixture'
        doc.core_properties.last_modified_by = ''
        doc.add_paragraph('Plain body')
        paragraph = getattr(doc.sections[0], slot).paragraphs[0]
        if omml:
            paragraph._p.append(math_run(17))
        else:
            paragraph.text = '$h+1$'
        doc.sections[0].different_first_page_header_footer = False
        doc.settings.odd_and_even_pages_header_footer = False
        doc.save(folder / f'{name}-input.docx')
        if name == 'review-first-header':
            # 只有其他区域增加目标时才处理首页页眉，是旧预检的漏转表现。
            doc.sections[0].footer.paragraphs[0].text = '$f+1$'
            doc.save(folder / 'review-header-control-input.docx')

    for preserved in (False, True):
        doc = Document()
        doc.core_properties.author = 'Synthetic test fixture'
        doc.core_properties.last_modified_by = ''
        doc.styles['Normal'].font.size = Pt(12)
        doc.add_paragraph('Before $a+1$ tail')
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].text = '$x' + '+1' * 45 + '$ / '
        cell.paragraphs[0]._p.append(math_run(18))
        cell.add_table(rows=1, cols=1).cell(0, 0).text = '$n+1$'
        paragraph = doc.add_paragraph('Between ')
        paragraph._p.append(math_run(19))
        paragraph.add_run(' / $b+2$ tail')
        doc.add_paragraph(r'\[c+3\]')
        if preserved:
            doc.add_paragraph(r'Keep $\text{测试}$ and reject ${oops$')
        # 第一张表格缩短，第二张表格增长，不能将两个累计偏移再次求和。
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '$z$'
        doc.add_paragraph('After $d+4$ tail')
        doc.add_paragraph()._p.append(math_run(20, True))
        name = 'review-preserved-input.docx' if preserved else 'review-positions-input.docx'
        doc.save(folder / name)


def check_review(folder, case):
    """通过持久化 XML 和独立 Word 重扫验证，禁止只对比转换器内部计算值。"""
    result = json.loads((folder / f'{case}.json').read_text(encoding='utf-8-sig'))
    if 'positions' in case or case == 'review-preserved':
        scan = json.loads((folder / f'{case}-scan.json').read_text(encoding='utf-8-sig'))
        body = [item for item in result['Items'] if not item.get('Experimental')]
        assert len(body) == len(scan['Items']) == (8 if case == 'review-preserved' else 6)
        for reported, actual in zip(body, scan['Items']):
            assert (reported['FinalStart'], reported['FinalEnd']) == (actual['Start'], actual['End']), (case, reported, actual)
            if reported['Status'] == 'SUCCEEDED' and result.get('Direction') != 'REVERSE':
                assert actual['Delimiter'] == 'MATH_TYPE'
            else:
                expected = reported['OutputText'] if reported['Status'] == 'SUCCEEDED' else reported['Source']
                assert actual['Source'] == expected, (case, reported, actual)
        assert {item['Kind'] for item in body} >= ({'MATH_TYPE'} if result.get('Direction') == 'REVERSE' else {'LATEX', 'OMML'})
        if case == 'review-preserved':
            assert result['FailedPreserved'] == result['RejectedPreserved'] == 1
    else:
        parts = document_parts(folder / f'{case}.docx')
        roots = [root for name, root in parts.items() if name.startswith(('header:', 'footer:'))]
        assert len(roots) == (2 if case == 'review-header-control' else 1)
        reverse = case.endswith('-reverse')
        assert sum(len(root.xpath("//*[local-name()='OLEObject' and @ProgID='Equation.DSMT4']")) for root in roots) == (0 if reverse else len(roots))
        assert not any(root.xpath('//m:oMath', namespaces=NS) for root in roots)
        for root in roots:
            value = ''.join(root.xpath('//w:t/text()', namespaces=NS))
            if reverse:
                assert ('$' in value or r'\[' in value), (case, value)
            else:
                assert '$' not in value and r'\[' not in value, (case, value)
        doc = Document(folder / f'{case}.docx')
        assert not doc.sections[0].different_first_page_header_footer
        assert not doc.settings.odd_and_even_pages_header_footer
    assert result['OriginalFileBytesUnchanged'] and result['WordReopenVerified']
    print(f'REVIEW_REGRESSION_PASSED: {case}')


def body_fixture(folder):
    """正文交错公式与隔离位置，独立覆盖顺序编号、字号和失败恢复。"""
    doc = Document()
    doc.core_properties.author = "Synthetic test fixture"
    doc.core_properties.last_modified_by = ""
    doc.styles["Normal"].font.size = Pt(14)
    doc.add_paragraph("Body start remains.")
    for value, tex in ((101, r"\[a+1\]"), (102, r"\[b+2\]")):
        p = doc.add_paragraph()
        p._p.append(math_run(value, True))
        doc.add_paragraph(tex)
    p = doc.add_paragraph("Inline $c+3$ / ")
    p._p.append(math_run(103))
    p.add_run(" end").bold = True
    table = doc.add_table(rows=1, cols=1)
    p = table.cell(0, 0).paragraphs[0]
    p.text = r"\[t+1\]"
    table.cell(0, 0).add_paragraph()._p.append(math_run(201, True))
    fill(doc.sections[0].header.paragraphs[0], "Header", 31)
    # 内容控件里的原生公式保持，避免默认主线越过既有安全边界。
    control = OxmlElement("w:sdt")
    properties = etree.SubElement(control, qn("w:sdtPr"))
    etree.SubElement(properties, qn("w:id"), {qn("w:val"): "3001"})
    content = etree.SubElement(control, qn("w:sdtContent"))
    etree.SubElement(content, qn("w:p")).append(math_run(301))
    doc._element.body.insert(len(doc._element.body) - 1, control)
    doc.add_paragraph("Body end remains.")
    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.5
    doc.save(folder / "body-input.docx")
    # 开发自检在第一个 OMML 已实际替换后注入失败。
    first = doc._element.body.xpath("./w:p//m:oMath")[0]
    first.find(f".//{{{M}}}t").text = "987654321"
    doc.save(folder / "body-failure-input.docx")
    first.find(f".//{{{M}}}t").text = "x+101"
    inline = doc._element.body.xpath("./w:p//m:oMath")[-1]
    inline.find(f".//{{{M}}}t").text = "987654321"
    doc.save(folder / "body-inline-failure-input.docx")


def failure_fixture(folder):
    """第二张表格用于证明第一张表失败不会跳过全部正文表格。"""
    doc = Document(folder / "input.docx")
    fill(doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0], "SecondTable", 301)
    doc.save(folder / "failure-input.docx")
    header_only_fixture(folder)
    body_fixture(folder)


def long_fixture(folder):
    """长原生公式包含 UTF-16 代理对，日志截断不能让正常区域转换失败。"""
    with ZipFile(folder / "input.docx") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name, data in list(parts.items()):
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        root = etree.fromstring(data)
        for math in root.xpath("//m:oMath", namespaces=NS):
            math.clear()
            for value in ("𝑥", "=") * 100 + ("1",):
                run = etree.SubElement(math, f"{{{M}}}r")
                etree.SubElement(run, f"{{{M}}}t").text = value
        parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(folder / "long-input.docx", "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def check_long(folder):
    result = json.loads((folder / "long-all.json").read_text(encoding="utf-8-sig"))
    assert result["Status"] == "COMPLETE" and result["Succeeded"] == 21 and not result["Issues"], result
    assert result["Experimental"]["Succeeded"] == 18
    actual = inventory(folder / "long-all.docx")
    assert sum(part["MathType"] for part in actual.values()) == 21
    assert sum(part["OMML"] for part in actual.values()) == 0
    # 证明样本确实命中了代理对被截断的边界，而非只测试了一条普通长字符串。
    sources = [item["Source"].encode("utf-16-le") for item in result["Experimental"]["Items"] if item["Kind"] == "OMML"]
    assert any(len(source) > 320 and 0xD800 <= int.from_bytes(source[318:320], "little") <= 0xDBFF for source in sources)
    excerpts = [bytes.fromhex(line.split()[3]).decode("utf-8") for line in (folder / "long-all.log").read_text(encoding="utf-8").splitlines() if line.startswith("UI_EXPERIMENTAL_ITEM ")]
    assert any(excerpt.endswith("…") for excerpt in excerpts)
    print("EXPERIMENTAL_LONG_UNICODE_DIAGNOSTIC_PASSED")


def region_element(parts, key):
    if key.startswith("body:table:"):
        return parts["word/document.xml"].xpath("/w:document/w:body/w:tbl", namespaces=NS)[int(key.rsplit(":", 1)[1]) - 1]
    if key.startswith("body:box:"):
        return parts["word/document.xml"].xpath("//w:txbxContent", namespaces=NS)[0]
    if key.startswith(("Headers:", "Footers:")):
        kind, section, variant = key.split(":")
        return parts[f"{'header' if kind == 'Headers' else 'footer'}:{section}:{dict({'1': 'default', '2': 'first', '3': 'even'})[variant]}"]
    kind, index = key.split(":")
    name = "footnote" if kind == "Footnotes" else "endnote"
    return parts[f"word/{name}s.xml"].xpath(f"/w:{name}s/w:{name}[@w:id='{index}']", namespaces=NS)[0]


def check_failure(folder, case, key, stage, reverse=False):
    """独立读取输出，核对成功公式、失败原式、容器与部分完成计数。"""
    result = json.loads((folder / f"{case}.json").read_text(encoding="utf-8-sig"))
    assert result["Status"] == "PARTIAL" and result["OriginalFileBytesUnchanged"] and result["WordReopenVerified"], result
    assert result["Targets"] == result["Succeeded"] + result["FailedPreserved"] + result["NotAttempted"], result
    baseline = document_parts(folder / ("failure-all.docx" if reverse else "failure-basic.docx"))
    output = document_parts(folder / f"{case}.docx")
    original = region_element(baseline, key)
    actual = region_element(output, key)
    def tex(root):
        return TEX.findall("".join(root.xpath(".//w:t/text()", namespaces=NS)))
    def maths(root):
        return ["".join(node.xpath(".//m:t/text()", namespaces=NS)) for node in root.xpath(".//m:oMath", namespaces=NS)]
    def objects(root):
        return len(root.xpath(".//*[local-name()='OLEObject' and @ProgID='Equation.DSMT4']"))
    total = target_count(folder / "failure-input.docx")
    assert total == 23
    if reverse:
        skipped = objects(original)
        assert result["Succeeded"] == total - skipped and objects(actual) == skipped, result
        assert sum(part["MathType"] for part in inventory(folder / f"{case}.docx").values()) == skipped
    else:
        skipped = len(maths(original)) if stage == "OMML" else len(tex(original)) + len(maths(original))
        assert result["Succeeded"] == total - skipped, result
        assert maths(actual) == maths(original), (key, stage, maths(actual), maths(original))
        if stage == "OMML":
            assert objects(actual) == len(tex(original)) and not tex(actual)
            assert result["FailedPreserved"] == skipped and not result["Issues"], result
        else:
            assert tex(actual) == tex(original) and not objects(actual), (key, stage)
            assert any(issue["Key"] == key for issue in result["Issues"]), result
            if stage == "RECOVERY":
                assert result["NotAttempted"] == skipped, result
        # 第二张表始终成功，验证表格失败的隔离边界。
        assert objects(region_element(output, "body:table:2")) == 2
    # 所有普通文字和表格属性保持，不能仅靠公式数量宣称成功。
    for name, before in baseline.items():
        after = output[name]
        for label in ("Basic inline", "Body", "Cell0", "Cell1", "Cell2", "Cell3", "Header", "Footer", "Textbox", "footnote", "endnote", "SecondTable"):
            if label in "".join(before.xpath("//w:t/text()", namespaces=NS)):
                assert label in "".join(after.xpath("//w:t/text()", namespaces=NS)), (name, label)
        for tag in ("tblPr", "trPr", "tcPr"):
            assert [etree.tostring(node, method="c14n", exclusive=True) for node in before.xpath(f"//w:{tag}", namespaces=NS)] == [etree.tostring(node, method="c14n", exclusive=True) for node in after.xpath(f"//w:{tag}", namespaces=NS)]
    print(f"PARTIAL_REGION_PASSED: {case}")


def check_body_failure(folder, case):
    from native_number_text import native_visible_numbers

    result = json.loads((folder / f"{case}.json").read_text(encoding="utf-8-sig"))
    assert result["Status"] == "PARTIAL" and result["OriginalFileBytesUnchanged"] and result["WordReopenVerified"], result
    parts = document_parts(folder / f"{case}.docx")
    body = parts["word/document.xml"]
    if case in ("partial-prepare", "partial-body-recovery", "partial-omml-recovery", "partial-inline-recovery"):
        if case == "partial-inline-recovery":
            assert result["Succeeded"] == 4 and result["NotAttempted"] == 2 and result["FailedPreserved"] == 0
            assert len(native_visible_numbers(body)) == 4
        elif case != "partial-prepare":
            assert result["Succeeded"] == 5 and result["NotAttempted"] == 1 and result["FailedPreserved"] == 0
            assert result["Issues"][0]["Stage"] == "RECOVERY"
            assert len(native_visible_numbers(body)) == 3
        else:
            assert result["Succeeded"] == 5 and result["FailedPreserved"] == 1
            assert len(native_visible_numbers(body)) == 3
        if case in ("partial-prepare", "partial-body-recovery"):
            assert r"\[a+1\]" in "".join(body.xpath("//w:t/text()", namespaces=NS))
    elif case == "partial-body-reverse":
        assert result["Succeeded"] == 1 and result["NotAttempted"] == 1 and result["FailedPreserved"] == 0
        assert len(body.xpath("//*[local-name()='OLEObject' and @ProgID='Equation.DSMT4']")) == 1
        assert native_visible_numbers(body) == ["(1)", "(2)"]
    elif case == "partial-numbering":
        assert result["Succeeded"] == 6 and result["NumbersAdded"] == 0 and result["NumbersUpdated"] == 0
        assert not native_visible_numbers(body) and result["Issues"][0]["Stage"] == "NUMBERING"
    elif case == "partial-display-omml":
        assert result["Succeeded"] == 9 and result["FailedPreserved"] == 1 and not result["Issues"], result
        table = body.xpath("/w:document/w:body/w:tbl", namespaces=NS)[0]
        assert len(table.xpath(".//m:oMathPara/m:oMath", namespaces=NS)) == 1
        assert len(table.xpath(".//*[local-name()='OLEObject']")) == 1
        assert len(native_visible_numbers(body)) == 4 and not native_visible_numbers(table)
    else:
        assert result["Succeeded"] == 0 and result["FailedPreserved"] == 0
        assert result["Issues"][0]["Stage"] == ("FORMATTING" if case == "partial-format" else "NUMBERING")
        source = document_parts(folder / "failure-numbered.docx")["word/document.xml"]
        assert native_visible_numbers(body) == native_visible_numbers(source)
        if case == "partial-format":
            assert result["ExistingDisplaysFormatted"] == 1
            assert etree.tostring(body.xpath("/w:document/w:body/w:p", namespaces=NS)[0].find(qn("w:pPr")), method="c14n", exclusive=True) == etree.tostring(source.xpath("/w:document/w:body/w:p", namespaces=NS)[0].find(qn("w:pPr")), method="c14n", exclusive=True)
    assert result["Targets"] == result["Succeeded"] + result["FailedPreserved"] + result["NotAttempted"]
    print(f"PARTIAL_BODY_PASSED: {case}")


def body_mixed(source, output):
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    body = root.find(qn("w:body"))
    first = etree.Element(qn("w:p"))
    first.append(math_run(401, True))
    body.insert(1, first)
    last = etree.Element(qn("w:p"))
    etree.SubElement(etree.SubElement(last, qn("w:r")), qn("w:t")).text = r"\[q+9\]"
    body.insert(len(body) - 1, last)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)


def check_body(folder, case):
    from native_number_text import native_visible_numbers
    from test_native_display import effective

    result = json.loads((folder / f"{case}.json").read_text(encoding="utf-8-sig"))
    source = Document(folder / "body-input.docx")
    doc = Document(folder / f"{case}.docx")
    root = document_parts(folder / f"{case}.docx")["word/document.xml"]
    expected_success = {"body-plain": 6, "body-numbered": 6, "body-repeat": 0, "body-mixed": 2, "body-inline": 2, "body-display": 4, "body-failure": 5, "body-inline-failure": 5, "body-extra": 10}[case]
    assert result["Succeeded"] == expected_success, result
    assert result["Status"] == ("PARTIAL" if case in ("body-failure", "body-inline-failure") else "COMPLETE"), result
    numbered = case != "body-plain" and case != "body-inline"
    count = 3 if case == "body-failure" else 6 if case == "body-mixed" else 4
    numbers = native_visible_numbers(root)
    assert numbers == [f"({i})" for i in range(1, count + 1)] if numbered else not numbers, numbers
    for p in doc.paragraphs:
        objects = p._p.xpath(".//*[local-name()='OLEObject']")
        if not objects:
            continue
        display = bool(native_visible_numbers(p._p)) or p.style.name == "MTDisplayEquation"
        if display:
            assert p.style.name == "MTDisplayEquation"
            assert len(p._p.xpath(".//w:tab")) == 2
        if case in ("body-numbered", "body-extra"):
            assert effective(p, "line_spacing") == (2.0 if display else Pt(30))
    if case in ("body-numbered", "body-extra"):
        assert all(float(item["NativeFullSize"]) == 18 for item in result["Items"] if item["Status"] == "SUCCEEDED" and not item.get("Experimental")), result
    if case == "body-repeat":
        assert result["NumbersAdded"] == 0 and result["NumbersUpdated"] == 4
    if case == "body-mixed":
        assert result["NumbersAdded"] == 2 and result["NumbersUpdated"] == 4
    if case in ("body-failure", "body-inline-failure"):
        assert result["FailedPreserved"] == 1
        assert "987654321" in "".join(root.xpath("//m:t/text()", namespaces=NS))
    # 额外位置中的公式没有编号，也不套用正文段落属性。
    for table in doc.tables:
        assert not native_visible_numbers(table._tbl)
        for before, after in zip(source.tables[0].cell(0, 0).paragraphs, table.cell(0, 0).paragraphs):
            assert before.paragraph_format.line_spacing == after.paragraph_format.line_spacing
    assert not native_visible_numbers(doc.sections[0].header._element)
    source_header = source.sections[0].header.paragraphs[0]
    saved_header = doc.sections[0].header.paragraphs[0]
    assert source_header.paragraph_format.line_spacing == saved_header.paragraph_format.line_spacing
    assert len(root.xpath("/w:document/w:body/w:sdt//m:oMath", namespaces=NS)) == 1
    assert "Body start remains." in "".join(root.xpath("//w:t/text()", namespaces=NS))
    assert root.xpath("//w:r[w:t=' end']/w:rPr/w:b", namespaces=NS), "inline suffix formatting lost"
    print(f"UNIFIED_BODY_PASSED: {case}")


if __name__ == "__main__":
    command, target = sys.argv[1:3]
    if command == "fixture":
        fixture(Path(target))
    elif command == "inventory":
        print(json.dumps(inventory(target), ensure_ascii=False))
    elif command == "check":
        check(target, sys.argv[3], sys.argv[4], sys.argv[5])
    elif command == "targets":
        print(target_count(target))
    elif command == "check-reverse":
        check_reverse(target, sys.argv[3])
    elif command == "body-fixture":
        body_fixture(Path(target))
    elif command == "body-mixed":
        body_mixed(target, sys.argv[3])
    elif command == "check-body":
        check_body(Path(target), sys.argv[3])
    elif command == "header-fixture":
        header_only_fixture(Path(target))
    elif command == "check-header":
        check_header_only(Path(target), sys.argv[3])
    elif command == 'review-fixture':
        review_fixture(Path(target))
    elif command == 'check-review':
        check_review(Path(target), sys.argv[3])
    elif command == "failure-fixture":
        failure_fixture(Path(target))
    elif command == "check-failure":
        check_failure(Path(target), sys.argv[3], sys.argv[4], sys.argv[5], len(sys.argv) > 6)
    elif command == "check-body-failure":
        check_body_failure(Path(target), sys.argv[3])
    elif command == "long-fixture":
        long_fixture(Path(target))
    elif command == "check-long":
        check_long(Path(target))
    elif command == "no-targets-fixture":
        no_targets_fixture(Path(target))
    elif command == "check-no-targets":
        check_no_targets(Path(target))
    elif command == "no-changes-fixture":
        no_changes_fixture(Path(target))
    elif command == "check-no-changes":
        check_no_changes(Path(target))
    else:
        raise SystemExit("unknown command")
