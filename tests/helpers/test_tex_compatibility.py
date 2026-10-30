"""核对兼容命令、嵌套多行和实验区域的实际页面及反向导出，未知命令仍可保留。"""
import json
import re
import sys
from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml
import fitz
from audit_docx import audit

# 只列已经用原生 TeX Input 对照过的写法，期望值描述数学内容而不限定翻译器排版。
CASES = [
    (r'\bm{x_1+\alpha}', ('x_1', r'\alpha')),
    (r'\textnormal{normal}', ('normal',)),
    (r'\textup{upright}', ('upright',)),
    (r'\textmd{medium}', ('medium',)),
    (r'\mathup{abc}+\alpha', ('abc', r'\alpha')),
    (r'\lvert x_2\rvert', ('x_2', '|')),
    (r'\mleft(\frac{x_3}{y_3}\mright)', ('x_3', 'y_3')),
    (r'a_1,\dotsc,a_n', ('a_1', 'a_n', r'\ldots')),
    (r'b_1\dotsm b_n', ('b_1', 'b_n', r'\cdots')),
    (r'f(c_1,\dotso,c_n)', ('c_1', 'c_n', r'\ldots')),
    (r'u\coloneqq v', ('u', ':', '=', 'v')),
    (r'w\eqqcolon z', ('w', '=', ':', 'z')),
    (r'\begin{align*}A_1&=B_1,&C_1&=D_1\\A_2&=B_2,&C_2&=D_2\end{align*}', tuple(f'{letter}_{i}' for i in (1, 2) for letter in 'ABCD')),
    (r'\begin{split}s_1&=t_1\\s_2&=t_2\end{split}', ('s_1', 't_1', 's_2', 't_2')),
    (r'\begin{gathered}[t]G_1=H_1\\G_2=H_2\end{gathered}', ('G_1', 'H_1', 'G_2', 'H_2')),
    (r'\begin{gather*}J_1=K_1\\J_2=K_2\end{gather*}', ('J_1', 'K_1', 'J_2', 'K_2')),
    (r'\begin{alignedat}{2}a_3&=b_3&c_3&=d_3\\e_3&=f_3&g_3&=h_3\end{alignedat}', tuple(letter+'_3' for letter in 'abcdefgh')),
    (r'\begin{alignat*}{2}a_4&=b_4&c_4&=d_4\end{alignat*}', tuple(letter+'_4' for letter in 'abcd')),
    (r'\sum_{\begin{subarray}{l}i<n\\j<m\end{subarray}} a_{ij}', ('i', 'n', 'j', 'm', 'a_ij')),
    (r'\begin{align*}h&=\left(\begin{gathered}\bm{p_1}\\\mathup{p_2}\end{gathered}\right)&r&=s\end{align*}', ('h', 'p_1', 'p_2', 'r', 's')),
]


def fixture(folder):
    document = Document()
    document.core_properties.author = 'Synthetic test fixture'
    document.core_properties.last_modified_by = ''
    document.styles['Normal'].font.size = Pt(12)
    for source, _ in CASES:
        document.add_paragraph(r'\['+source+r'\]')
    document.add_paragraph(r'Inline $\bm{v_1}+\lvert v_2\rvert$ stays inline.')
    # 单行 alignedat 源码跨段但没有 TeX 分行符，也必须按结构化公式识别。
    document.add_paragraph(r'\[\begin{alignedat}{1}q_1')
    document.add_paragraph(r'&=q_2\end{alignedat}\]')
    table = document.add_table(rows=len(CASES), cols=1)
    for cell, (source, _) in zip((row.cells[0] for row in table.rows), CASES):
        cell.text = r'\['+source+r'\]'
    document.add_paragraph('End stays exact.')
    document.save(folder / 'tex-compatibility-input.docx')


def check(folder):
    target = folder / 'tex-compatibility.docx'
    report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
    reverse = json.loads((folder / 'tex-compatibility-reverse.json').read_text(encoding='utf-8-sig'))
    count = 2*len(CASES)+2
    assert report['Status'] == reverse['Status'] == 'COMPLETE'
    assert report['Succeeded'] == reverse['Succeeded'] == count
    assert report['Experimental']['Succeeded'] == len(CASES)
    assert report['OriginalFileBytesUnchanged'] and report['WordReopenVerified']
    assert audit(target)['MathType'] == count and audit(folder / 'tex-compatibility-reverse.docx')['MathType'] == 0
    # 报告原文仍使用用户写法；兼容改写只影响 MathType 的导入副本。
    sources = Counter(item['Source'] for item in report['Items'])
    for source, _ in CASES:
        assert sources[r'\['+source+r'\]'] == 2
    body = [item for item in reverse['Items'] if not item.get('Experimental')]
    assert r'\mathbf' in body[0]['OutputText'] or r'\boldsymbol' in body[0]['OutputText']
    for item, (_, tokens) in zip(body, CASES):
        # 厂商会把粗体分别写在下标等子项上，先去字体命令再核对数学内容。
        tex = re.sub(r'\\(?:mathbf|mathrm|mathit|boldsymbol|text|rm)\b', '', item['OutputText'])
        tex = re.sub(r'\s|[{}]', '', tex)
        for token in tokens:
            assert token in tex, (token, tex)
    assert all(token in body[-2]['OutputText'] for token in ('v', '1', '2'))
    assert all(token in body[-1]['OutputText'] for token in ('q', '1', '2'))
    assert Document(folder / 'tex-compatibility-reverse.docx').paragraphs[-1].text == 'End stays exact.'
    # 页面检查只在开发回归中执行，生产转换不增加逐式导出或 PDF 核验。
    with fitz.open(str(target)+'.pdf') as pdf:
        for page in pdf:
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        color = span['color']
                        assert not ((color >> 16 & 255) > 153 and (color >> 8 & 255) < 77 and (color & 255) < 77), span
            for drawing in page.get_drawings():
                for color in (drawing['color'], drawing['fill']):
                    assert color is None or not (color[0] > .6 and color[1] < .3 and color[2] < .3), drawing
    print(json.dumps({'Status': 'PASSED', 'MathType': count, 'NoRedResiduals': True, 'OriginalSyntaxPreserved': True, 'NestedAndExperimental': True}))


# 顶饰单独集中回归，不必为了验证本次修改运行整个项目的测试。
ACCENT_TEX = [
    ('Ŷ+Y\u0302', (r'\hat', 'Y')),
    (r'\mathrm{Ȳ}_i^2', (r'\bar', 'Y', 'i', '2')),
    ('Ȭ', (r'\bar', r'\tilde', 'O')),
    (r'\widetilde{XY}', (r'\widetilde', 'X', 'Y')),
    (r'\mathring{Y}', (r'\circ', 'Y')),
    (r'\hat{Y}+\dot{Y}+\vec{Y}', (r'\hat', r'\dot', r'\vec')),
]
ACCENT_OMML = [
    ('precomposed', '<m:r><m:t>Ŷ</m:t></m:r>', (r'\hat', 'Y')),
    ('word-middle', '<m:r><m:t>XŶZ</m:t></m:r>', ('X', r'\hat', 'Y', 'Z')),
    ('combining', '<m:r><m:t>XY\u0302</m:t></m:r>', ('X', r'\hat', 'Y')),
    ('split-run', '<m:r><m:t>Y</m:t></m:r><m:r><m:t>\u0302</m:t></m:r>', (r'\hat', 'Y')),
    ('stacked', '<m:r><m:t>Ȭ</m:t></m:r>', (r'\overline', r'\tilde', 'O')),
    ('bold', '<m:r><m:rPr><m:sty m:val="b"/></m:rPr><m:t>Ŷ</m:t></m:r>', (r'\hat', r'\mathbf', 'Y')),
]
for label, code, expected in [('hat', '0302', r'\hat'), ('bar', '0304', r'\bar'),
                              ('breve', '0306', r'\smile'), ('dot', '0307', r'\dot'),
                              ('tilde', '0303', r'\tilde')]:
    ACCENT_OMML.append((label, f'<m:acc><m:accPr><m:chr m:val="&#x{code};"/></m:accPr><m:e><m:r><m:t>Y</m:t></m:r></m:e></m:acc>', (expected, 'Y')))


def accent_fixture(folder):
    document = Document()
    document.core_properties.author = 'Synthetic test fixture'
    document.core_properties.last_modified_by = ''
    document.styles['Normal'].font.size = Pt(12)
    namespace = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    for source, _ in ACCENT_TEX:
        document.add_paragraph('LaTeX $' + source + '$')
    for label, body, _ in ACCENT_OMML:
        paragraph = document.add_paragraph('Word ' + label + ' ')
        paragraph._p.append(parse_xml(f'<m:oMath xmlns:m="{namespace}">{body}</m:oMath>'))
    # 表格同时覆盖 Unicode LaTeX 和 Word 原生公式两个共用入口。
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = '$Ŷ+Y\u0302$'
    table.cell(1, 0).paragraphs[0]._p.append(parse_xml(f'<m:oMath xmlns:m="{namespace}"><m:r><m:t>Ŷ</m:t></m:r></m:oMath>'))
    document.add_paragraph(r'Unmapped $\text{中文}$ stays unchanged.')
    document.add_paragraph('Unmapped $Ç$ stays unchanged.')
    document.save(folder / 'accents-input.docx')


def accent_check(folder):
    report = json.loads((folder / 'accents.json').read_text(encoding='utf-8-sig'))
    count = len(ACCENT_TEX) + len(ACCENT_OMML) + 2
    assert report['Status'] == 'PARTIAL' and report['Succeeded'] == count
    assert report['Experimental']['Succeeded'] == 2
    failures = [item for item in report['Items'] if item['Status'] == 'FAILED_PRESERVED']
    assert len(failures) == 2 and all(item['Reason'] == 'UNSUPPORTED_NON_ASCII_TEX' for item in failures)
    assert report['OriginalFileBytesUnchanged'] and report['WordReopenVerified']
    assert audit(folder / 'accents.docx')['MathType'] == count
    body = (folder / 'accents-native.txt').read_text(encoding='utf-8').splitlines()
    assert len(body) == count
    expected = [tokens for _, tokens in ACCENT_TEX] + [tokens for _, _, tokens in ACCENT_OMML]
    for tex, tokens in zip(body, expected):
        for token in tokens:
            assert token in tex, (token, tex)
    assert all(r'\hat' in tex for tex in body[-2:])
    # 这里只做开发页面检查；正式转换继续沿用既有缓存回显和成功判定。
    with fitz.open(str(folder / 'accents.docx') + '.pdf') as pdf:
        page_text = ''.join(page.get_text() for page in pdf)
        assert '\ufffd' not in page_text and '\u25c6' not in page_text
    print(json.dumps({'Status': 'PASSED', 'Converted': count, 'Preserved': 2, 'BodyAndTable': True}))


def accent_reverse_check(folder):
    # 短音符的嵌套美元符只影响该对象，其余公式必须正常交付。
    target = folder / 'accents-reverse.docx'
    report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
    count = len(ACCENT_TEX) + len(ACCENT_OMML) + 2
    assert report['Status'] == 'PARTIAL' and report['FailedPreserved'] == 1
    assert report['Succeeded'] == count + 1 and report['Experimental']['Succeeded'] == 2
    failures = [item for item in report['Items'] if item['Status'] == 'FAILED_PRESERVED']
    assert failures[0]['Reason'] == 'UNSUPPORTED_OR_INVALID_NATIVE_EXPORT'
    assert report['OriginalFileBytesUnchanged'] and report['WordReopenVerified']
    assert audit(target)['MathType'] == 1
    print('ACCENT_REVERSE_PARTIAL_PASSED')


if __name__ == '__main__':
    mode, directory = sys.argv[1:3]
    {'fixture': fixture, 'check': check, 'accent-fixture': accent_fixture, 'accent-check': accent_check,
     'accent-reverse-check': accent_reverse_check}[mode](Path(directory))
