"""以 Word PDF 的可见编号独立核对原生删除、更新、格式化、追加和重开。"""
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import fitz
from docx import Document
from docx.shared import Pt

folder = Path(sys.argv[2])
if sys.argv[1] == 'fixture':
    doc = Document()
    doc.core_properties.author = "Synthetic test fixture"
    doc.core_properties.last_modified_by = ""
    doc.styles['Normal'].font.size = Pt(12)
    doc.add_paragraph('前文保持。')
    for i in range(1, 6):
        doc.add_paragraph(r'\[x_{' + str(i) + r'}+1\]')
    doc.add_paragraph('行内 $a+b$ 保持同行。')
    doc.add_paragraph('尾部正文保持。')
    doc.save(folder / 'input.docx')
    print('NATIVE_NUMBER_FIXTURE_CREATED')
    sys.exit()

count = int(sys.argv[3])
custom = sys.argv[4] == 'custom'
start = 27 if custom else 1
if (folder / 'prepended.pdf').exists():
    with fitz.open(folder / 'prepended.pdf') as pdf:
        actual = re.findall(r'\((\d+)\)', ''.join(page.get_text() for page in pdf))
    # mixed 原件含厂商章节断点；断点前新增 1，断点后的原序列仍从 1 开始。
    expected = (['1'] + list(map(str, range(1, count + 1)))) if sys.argv[4] == 'mixed' else list(map(str, range(1, count + 2)))
    assert actual == expected, ('prepended', actual, expected)
    print(json.dumps({'Stage': 'prepended', 'Status': 'PASSED', 'VisibleNumbers': actual}))
# 转换和原生命令执行后均检查真实 DOCX，不能用 PDF 不显示隐藏文字来冒充无断点。
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
for filename in ('converted.docx', 'native-checked.docx'):
    with ZipFile(folder / filename) as package:
        root = etree.fromstring(package.read('word/document.xml'))
        styles = etree.fromstring(package.read('word/styles.xml'))
    sections = root.xpath('//w:instrText[contains(., "MACROBUTTON MTEditEquationSection")]', namespaces=ns)
    expected_sections = 1 if sys.argv[4] == 'mixed' else 0
    assert len(sections) == expected_sections, (filename, '程序或原生命令新增了公式断点', len(sections))
    instructions = ''.join(root.xpath('//w:instrText/text()', namespaces=ns))
    if expected_sections:
        assert r'SEQ MTChap \r 5' in instructions and r'SEQ MTSec \r 7' in instructions, '默认转换重置了原生章、节值'
    else:
        assert not styles.xpath('//w:style[w:name/@w:val="MTEquationSection"]', namespaces=ns), '新增了章节断点样式'
        assert '公式章' not in instructions and 'Equation Chapter' not in instructions, '新增了可见的章节文字'
    print(json.dumps({'File': filename, 'NoNewEquationBreaks': True, 'SectionCount': len(sections)}))
for name, total, bracket in (
    ('original', count, False), ('deleted-updated', count - 1, False),
    ('formatted', count - 1, True), ('appended', count, True),
    ('reformatted', count, False), ('reopened', count, False),
):
    with fitz.open(folder / (name + '.pdf')) as pdf:
        text = ''.join(page.get_text() for page in pdf)
    if custom and name in ('original', 'deleted-updated'):
        actual = re.findall(r'\[IV.b-([A-Z]+)\]', text)
        expected = [chr(ord('A') + n - 27) * 2 for n in range(start, start + total)]
    else:
        actual = re.findall(r'\[(\d+)\]' if bracket else r'\((\d+)\)', text)
        expected = list(map(str, range(start, start + total)))
    assert actual == expected, (name, actual, text)
    assert not re.search(r'\(\d+\)' if bracket else r'\[\d+\]', text), (name, text)
    print(json.dumps({'Stage': name, 'Status': 'PASSED', 'VisibleNumbers': actual}))
