"""合成版式提醒和密集实验区域小样；核对报告及原有行距，不使用私人文档。"""
import json
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt
from lxml import etree

from test_experimental import fixture, math_run, inventory, document_parts, NS


def make_fixture(folder):
    fixture(folder)
    document = Document(folder / "input.docx")
    wide = document.add_paragraph(r"\[" + "+".join(f"x_{{{i}}}" for i in range(1, 36)) + r"\]")
    wide.paragraph_format.left_indent = Pt(18)
    wide.paragraph_format.right_indent = Pt(18)
    style = document.styles.add_style("SyntheticExactSpacing", WD_STYLE_TYPE.PARAGRAPH)
    style.paragraph_format.line_spacing = Pt(6)
    document.add_paragraph(r"Inherited $\frac{a+2}{b+3}$ tail", style)
    minimum = document.add_paragraph(r"Minimum $\frac{a+4}{b+5}$ tail")
    minimum.paragraph_format.line_spacing = Pt(6)
    minimum.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    # 同一页眉含多式，才能观察逐式整区快照的重复开销。
    header = document.sections[0].header
    for index in range(4):
        paragraph = header.add_paragraph(f"Dense{index} $x+{index + 201}$ / ")
        paragraph._p.append(math_run(index + 301))
        paragraph.paragraph_format.line_spacing = Pt(6)
    document.save(folder / "layout-input.docx")
    print("LAYOUT_FIXTURE_CREATED")


def check(folder, case, report_path=None, output_path=None):
    report = json.loads((report_path or folder / f"{case}.json").read_text(encoding="utf-8-sig"))
    assert report["Status"] == "COMPLETE", report["Issues"]
    assert report["FailedPreserved"] == report["RejectedPreserved"] == report["NotAttempted"] == 0
    warnings = report["Warnings"]
    target = output_path or folder / f"{case}.docx"
    contents = inventory(target)
    if case == "skip-omml":
        original = document_parts(folder / "layout-input.docx")
        actual = document_parts(target)
        # 独立核对保留公式的顺序和文字；生产核验另检查完整 XML 数学结构。
        for name, root in original.items():
            source_maths = ["".join(node.xpath(".//m:t/text()", namespaces=NS)) for node in root.xpath("//m:oMath", namespaces=NS)]
            target_maths = ["".join(node.xpath(".//m:t/text()", namespaces=NS)) for node in actual[name].xpath("//m:oMath", namespaces=NS)]
            assert source_maths == target_maths, name
        assert report["ConvertOmml"] is False
    else:
        assert sum(part["OMML"] for part in contents.values()) == 0
    assert sum(part["MathType"] for part in contents.values()) == (0 if case == "reverse" else report["Succeeded"])
    if case == "fast":
        assert report["LayoutWarningsChecked"] is False and warnings == []
        assert report["PostConversionScans"] == 0 and report["WordReopenVerified"] is False
    elif case == "reverse":
        assert report["LayoutWarningsChecked"] is False and warnings == []
    else:
        assert report["LayoutWarningsChecked"] is True
        codes = {item["Code"] for item in warnings}
        assert codes == {"FORMULA_TOO_WIDE", "EXACT_LINE_SPACING"}, warnings
        assert any("x_{35}" in item["Source"] and item["Code"] == "FORMULA_TOO_WIDE" for item in warnings)
        assert any("a+2" in item["Source"] and item["Code"] == "EXACT_LINE_SPACING" for item in warnings)
        assert not any("a+4" in item["Source"] for item in warnings)
        assert any(item["Region"] != "正文" and item["Code"] == "EXACT_LINE_SPACING" for item in warnings)
        assert all(item["SourceLocation"] and item["Actual"] > item["Limit"] + 1 for item in warnings)
        assert len({(item["Id"], item["Code"]) for item in warnings}) == len(warnings)
        # 版式提醒只读；页眉中人为设置的固定行距必须原样保留。
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with ZipFile(output_path or folder / f"{case}.docx") as archive:
            # Word 会重命名页眉部件，按样本文本找段落，不依赖部件编号。
            paragraphs = [paragraph for name in archive.namelist() if name.startswith("word/header") and name.endswith(".xml")
                          for paragraph in etree.fromstring(archive.read(name)).xpath("//w:p", namespaces=ns)
                          if "Dense" in "".join(paragraph.xpath("./w:r/w:t/text()", namespaces=ns))]
        values = [value for paragraph in paragraphs for value in paragraph.xpath("./w:pPr/w:spacing[@w:lineRule='exact']/@w:line", namespaces=ns)]
        assert len(values) == 4 and set(values) == {"120"}, values
    print(json.dumps({"Case": case, "Status": "PASSED", "Succeeded": report["Succeeded"], "Warnings": len(warnings)}, ensure_ascii=False))


if __name__ == "__main__":
    action, directory = sys.argv[1], Path(sys.argv[2])
    if action == "fixture":
        make_fixture(directory)
    else:
        check(directory, action, *(Path(value) for value in sys.argv[3:]))
