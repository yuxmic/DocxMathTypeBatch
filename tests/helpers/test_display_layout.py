"""验证真实 Word PDF 中的公式中心、右编号及保留的源段落格式；仅用于开发回归。"""
import json
import sys
from pathlib import Path
import fitz
from docx import Document
from audit_docx import audit


def effective(paragraph, name, default=0):
    # Word 保存后可能将直接格式归入样式，比较实际继承值。
    value = getattr(paragraph.paragraph_format, name)
    style = paragraph.style
    while value is None and style:
        value = getattr(style.paragraph_format, name)
        style = style.base_style
    return default if value is None else value


def points(value):
    return value.pt if hasattr(value, 'pt') else value


def check_layout():
    source, target = map(Path, sys.argv[1:3])
    numbered = len(sys.argv) > 3 and sys.argv[3] == '--numbers'
    before, after = Document(source), Document(target)
    report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
    assert report['Status'] == 'COMPLETE' and report['Succeeded'] == 16
    assert report['NumbersAdded'] == (8 if numbered else 0)
    sizes = [float(item['NativeFullSize']) for item in report['Items']]
    assert len(sizes) == 16 and sizes.count(14) == 1 and all(size in (12, 14) for size in sizes)
    assert audit(target)['MathType'] == 16
    assert len(before.paragraphs) == len(after.paragraphs)
    assert before.tables[0].cell(0, 0).text == after.tables[0].cell(0, 0).text
    display_sources = []
    for original, actual in zip(before.paragraphs, after.paragraphs):
        display = original.text.lstrip().startswith(r'\[')
        for name in ('space_before', 'space_after', 'left_indent', 'right_indent', 'line_spacing'):
            assert effective(original, name) == effective(actual, name), (original.text, name)
        if display:
            assert points(effective(actual, 'first_line_indent')) == 0, original.text
            assert len(actual._p.xpath('.//*[local-name()="OLEObject"]')) == 1
            display_sources.append(original)
        else:
            assert (original.alignment or 0) == (actual.alignment or 0)
            assert effective(original, 'first_line_indent') == effective(actual, 'first_line_indent')
            # 行内公式可替换，周围正文和普通段落内容保持。
            assert actual.text == original.text.replace('$a_1+1$', '')
    assert len(display_sources) == 8

    # 分数线中心可从厂商实际绘制结果测量，OLE 域起点的 Word 坐标不能代替对象位置。
    pdf = fitz.open(str(target) + '.pdf')
    fractions = []
    for page_no, page in enumerate(pdf):
        for drawing in page.get_drawings():
            rect = drawing['rect']
            if rect.height < .1 and 20 < rect.width < 40:
                fractions.append((page_no, (rect.x0 + rect.x1) / 2))
    assert len(fractions) == 8
    section = after.sections[0]
    positions, numbers = [], []
    for index, (original, (page_no, center)) in enumerate(zip(display_sources, fractions), 1):
        left = section.left_margin.pt + points(effective(original, 'left_indent'))
        right = section.page_width.pt - section.right_margin.pt - points(effective(original, 'right_indent'))
        offset = center - (left + right) / 2
        assert abs(offset) <= 1, (index, 'center', offset)
        positions.append(dict(Formula=index, Page=page_no + 1, OffsetPoints=offset))
        if numbered:
            hits = [rect for page in pdf for rect in page.search_for(f'({index})')]
            assert len(hits) == 1 and abs(hits[0].x1 - right) <= 1, (index, 'right number', hits)
            numbers.append(dict(Formula=index, OffsetPoints=hits[0].x1 - right))
    print(json.dumps(dict(Status='PASSED',Centered=8,Numbers=len(numbers),NativeSizesPreserved=True,
                         OriginalSpacingPreserved=True,SurroundingParagraphsPreserved=True,
                         TableUnchanged=True,Positions=positions,NumberPositions=numbers), ensure_ascii=False))


if __name__ == '__main__':
    check_layout()
