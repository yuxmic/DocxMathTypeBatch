"""独立读取 DOCX，核对行距、编号可见结果、原有段距和表格。"""
import json
import sys
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from native_number_text import native_visible_numbers
from test_native_display import effective

mode, folder = sys.argv[1], Path(sys.argv[2])
if mode == 'omml-switch-fixture':
    # 复用包含正文、表格、文本框、页眉页脚及注释的公开样本。
    from test_experimental import fixture, TEX, NS
    from lxml import etree

    fixture(folder)
    with ZipFile(folder / 'input.docx') as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name in parts:
        if name.startswith('word/') and name.endswith('.xml'):
            root = etree.fromstring(parts[name])
            for node in root.xpath('//w:t', namespaces=NS):
                if node.text:
                    node.text = TEX.sub('', node.text)
            if name == 'word/document.xml':
                # OMML 内部的定界符形文本仍是原生公式，不能被 LaTeX 扫描器越界处理。
                root.xpath('//m:oMath//m:t', namespaces=NS)[0].text = '$native_only$'
            parts[name] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(folder / 'omml-only-input.docx', 'w', ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    print('OMML_SWITCH_FIXTURE_PASSED')
    sys.exit()

if mode == 'omml-switch-check':
    from lxml import etree
    from test_experimental import document_parts, inventory, TEX, NS, M

    case = sys.argv[3]
    only = case == 'off-no-targets'
    enabled = case == 'default-on'
    reverse = case == 'off-reverse'
    original = folder / ('omml-only-input.docx' if only else 'input.docx')
    target = folder / (case + '.docx')
    report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
    before, after = document_parts(original), document_parts(target)
    inputs, outputs = inventory(original), inventory(target)
    latex_count = sum(len(TEX.findall(value['Text'])) for value in inputs.values())
    omml_count = sum(value['OMML'] for value in inputs.values())
    expected = latex_count + (omml_count if enabled else 0)
    assert report['Status'] == 'COMPLETE' and report['ConvertOmml'] is enabled, report
    assert report['Succeeded'] == report['Targets'] == expected, report
    assert not report['FailedPreserved'] and not report['RejectedPreserved'] and not report['NotAttempted']
    assert not report['Issues'] and not report['Experimental']['Issues'], report
    assert report['Experimental']['RejectedPreserved'] == report['Experimental']['FailedPreserved'] == 0
    if not enabled:
        assert all(item.get('Kind') != 'OMML' for item in report['Items'])
        assert all(item['Kind'] != 'OMML' for item in report['Experimental']['Items'])

    def native_equations(root):
        # 忽略 Word 的运行拆分和默认格式补齐，仍逐层比较数学结构、内容与行外位置。
        def structure(node):
            if not isinstance(node.tag, str) or not node.tag.startswith('{' + M + '}'):
                return []
            local = etree.QName(node).localname
            if local.endswith('Pr'):
                return []
            if local == 't':
                return [node.text or '']
            children = []
            for child in node:
                for item in structure(child):
                    if isinstance(item, str) and children and isinstance(children[-1], str):
                        children[-1] += item
                    else:
                        children.append(item)
            return children if local == 'r' else [(local, tuple(sorted(node.attrib.items())), tuple(children))]

        result = []
        for equation in root.xpath('//m:oMath', namespaces=NS):
            result.append((bool(equation.xpath('ancestor::m:oMathPara', namespaces=NS)),
                           structure(equation)))
        return result

    assert set(before) == set(after), (set(before), set(after))
    for name in before:
        expected_native = [] if enabled else native_equations(before[name])
        assert native_equations(after[name]) == expected_native, (case, name)
        expected_objects = 0 if reverse else len(TEX.findall(inputs[name]['Text'])) + (inputs[name]['OMML'] if enabled else 0)
        assert outputs[name]['MathType'] == expected_objects, (case, name, outputs[name])
    assert sum(value['OMML'] for value in outputs.values()) == (0 if enabled else omml_count)
    if only:
        log = target.with_suffix('.log').read_text(encoding='utf-8')
        assert not report['Items'] and not report['Experimental']['Items']
        assert report['Experimental']['Skipped'] == 'NO_TARGETS', report
        assert 'SHARED_MATHTYPE_PID=' not in log and 'DEFAULT_PREFERENCES_LENGTH' not in log, log
    print(json.dumps({'Status': 'PASSED', 'Case': case, 'ConvertOmml': enabled,
                      'Succeeded': expected, 'NativeEquationsPreserved': not enabled}, ensure_ascii=False))
    sys.exit()

if mode == 'fixture':
    doc = Document()
    doc.core_properties.author = "Synthetic test fixture"
    doc.core_properties.last_modified_by = ""
    doc.styles['Normal'].font.size = Pt(12)
    for text in ('普通正文不改。', '正文 $a+b$ 和 $c+d$ 同行。', r'\[\frac{x+1}{y+2}\]', '$$x+1$$', '结尾正文不改。'):
        p = doc.add_paragraph(text)
        f = p.paragraph_format
        f.line_spacing = Pt(20)
        f.space_before, f.space_after = Pt(6), Pt(9)
        f.left_indent, f.right_indent = Pt(12), Pt(6)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '$t_1$ 不转换'
    # Word 要求表格之后有末段，样本直接提供，避免把自动末段算作转换拆段。
    doc.paragraphs[-1]._p.addprevious(table._tbl)
    doc.save(folder / 'advanced-input.docx')
    print('ADVANCED_FIXTURE_PASSED')
    sys.exit()

source, target = Document(folder / 'advanced-input.docx'), Document(folder / (mode + '.docx'))
report = json.loads((folder / (mode + '.json')).read_text(encoding='utf-8-sig'))
assert report['Status'] == 'COMPLETE' and report['Succeeded'] == 4 and report['ConvertOmml'] is True
assert len(source.paragraphs) == len(target.paragraphs) == 5
assert target.tables[0].cell(0, 0).text == source.tables[0].cell(0, 0).text

def spacing(p):
    value = p._p.find('./' + qn('w:pPr') + '/' + qn('w:spacing'))
    # 原生显示段落会将相同属性归入样式，按有效属性比较而非只读直接格式。
    style = p.style
    while value is None and style is not None:
        value = style.element.find('./' + qn('w:pPr') + '/' + qn('w:spacing'))
        style = style.base_style
    return value.get(qn('w:lineRule'), 'auto'), int(value.get(qn('w:line')))

rules = {
    'follow': ('exact', 400), 'single': ('auto', 240), 'onehalf': ('auto', 360),
    'double': ('auto', 480), 'exact:24.000000': ('exact', 480),
    'atleast:28.000000': ('atLeast', 560), 'multiple:1.250000': ('auto', 300),
}
for index, (before, after) in enumerate(zip(source.paragraphs, target.paragraphs)):
    for name in ('space_before', 'space_after', 'left_indent', 'right_indent'):
        assert effective(before, name) == effective(after, name), (mode, index, name)
    if index in (0, 4):
        assert before.text == after.text and spacing(before) == spacing(after)
    else:
        setting = report['InlineSpacing' if index == 1 else 'DisplaySpacing']
        assert spacing(after) == rules[setting], (mode, index, spacing(after), setting)
        if setting != 'follow':
            assert after._p.find('./' + qn('w:pPr') + '/' + qn('w:snapToGrid')).get(qn('w:val')) == '0'

expected = {
    'alphabetic': ['[IV.b-AA]', '[IV.b-BB]'],
    'roman': ['ix', 'x'], 'numeric': ['（3-2-9）', '（3-2-10）'],
}
visible = ''.join(native_visible_numbers(target.element.body))
for number in expected.get(mode, []):
    assert number in visible, (number, visible)
assert target.paragraphs[1].text.startswith('正文 ') and target.paragraphs[1].text.endswith(' 同行。')
print(json.dumps({'Status': 'PASSED', 'Case': mode, 'ParagraphSpacingVerified': True,
                  'UnselectedContentUnchanged': True, 'NumberTextVerified': expected.get(mode, [])}, ensure_ascii=False))
