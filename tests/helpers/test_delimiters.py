"""仅使用合成文档验证定界符、上下文和破损原文保留。"""
import json
import sys
from pathlib import Path
from docx import Document
from audit_docx import audit

source, target = map(Path, sys.argv[1:3])
before, after = Document(source), Document(target)
report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
items = [item for item in report['Items'] if item['Status'] == 'SUCCEEDED']
assert len(items) == 12 and sum(item['Display'] for item in items) == 7
assert report['Status'] == 'PARTIAL' and report['RejectedPreserved'] == 2
assert audit(target)['MathType'] == 12
inline_sources = {r'$$a+1$$', r'\[b+1\]', r'\(c+1\)', '$a$'}
assert all(item['Display'] == (item['Source'] not in inline_sources) for item in items)
assert all(item['AddedParagraphs'] == 0 for item in items)
for index in (9, 11, 12, 13):
    assert before.paragraphs[index].text == after.paragraphs[index].text
assert len(before.paragraphs) == len(after.paragraphs)
print('GENERATED_DELIMITERS_PASSED')
