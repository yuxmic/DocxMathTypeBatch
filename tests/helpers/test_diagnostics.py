"""生成公开分页样本，独立核对诊断报告的位置、原式和保留状态。"""
import hashlib
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def fixture(folder):
    document = Document()
    document.core_properties.author = "Synthetic diagnostic fixture"
    document.core_properties.last_modified_by = ""
    document.styles["Normal"].font.size = Pt(12)
    document.styles["Normal"].paragraph_format.space_after = Pt(0)
    document.styles["Normal"].paragraph_format.line_spacing = 1
    section = document.sections[0]
    # 原文六个短公式能放在一页；72pt 转换后必须独立读输出确认后页真正移动。
    section.page_width = section.page_height = Pt(360)
    section.top_margin = section.bottom_margin = Pt(36)
    section.left_margin = section.right_margin = Pt(36)
    section.header_distance = section.footer_distance = Pt(12)
    section.header.paragraphs[0].text = "HEADER_ONE $x_{1$"
    for index in range(6):
        document.add_paragraph(r"\[x+" + str(index + 1) + r"\]")
    document.add_page_break()
    document.add_paragraph("DIAGNOSTIC_PAGE_TWO")
    document.add_paragraph("BODY_TWO $x_{1$")
    document.add_paragraph("FAILURE_TARGET $f+7$")
    failure_paragraph = len(document.paragraphs)
    for index, label in enumerate(("TABLE_ONE", "TABLE_TWO")):
        # 空段隔开相邻表格，避免 Word 将两张表合为同一张。
        if index:
            document.add_paragraph()
        document.add_table(rows=1, cols=1).cell(0, 0).text = label + " $x_{1$"
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = "HEADER_TWO $x_{1$"
    # 分节重启显示页码，不改变日志所用的连续物理页序（第三页）。
    numbering = OxmlElement("w:pgNumType")
    numbering.set(qn("w:start"), "1")
    section._sectPr.append(numbering)
    document.add_paragraph("DIAGNOSTIC_RESTART $x_{1$")
    source = folder / "pages-input.docx"
    document.save(source)
    metadata = {"FailureParagraph": failure_paragraph, "SourcePage": 2,
                "RestartPage": 3, "SourceSHA256": hashlib.sha256(source.read_bytes()).hexdigest()}
    (folder / "pages-fixture.json").write_text(json.dumps(metadata), encoding="utf-8")


def check(folder, reflow):
    source = folder / "input.docx"
    metadata = json.loads((folder.parent / "pages-fixture.json").read_text(encoding="utf-8"))
    assert hashlib.sha256(source.read_bytes()).hexdigest() == metadata["SourceSHA256"]
    result = json.loads((folder / "report-0.json").read_text(encoding="utf-8-sig"))
    ui = json.loads((folder / "batch-check.json").read_text(encoding="utf-8-sig"))["Items"][0]
    assert result["Status"] == "PARTIAL" and ui["State"] == 3, result
    assert result["Succeeded"] == (6 if reflow else 7), result
    assert result["FailedPreserved"] == (1 if reflow else 0) and result["RejectedPreserved"] == 6, result
    assert not result["NotAttempted"] and not result["Issues"] and not result["Experimental"]["Issues"], result
    assert result["OriginalFileBytesUnchanged"] and result["WordReopenVerified"]
    assert result["MathematicalEquivalenceVerified"] is False
    items = result["Items"]
    body = [item for item in items if not item.get("Experimental")]
    body_rejected = [item for item in body if item["Status"] == "REJECTED_PRESERVED"]
    assert len(body_rejected) == 2, body_rejected
    page_two, restart = body_rejected
    assert page_two["Page"] == 2 and "原文第 2 页" in page_two["SourceLocation"], page_two
    assert restart["Page"] == 3 and "原文第 3 页" in restart["SourceLocation"], restart
    failure = next(item for item in body if item["Source"] == "$f+7$")
    assert failure["Page"] == 2 and failure["Status"] == ("FAILED_PRESERVED" if reflow else "SUCCEEDED"), failure
    rejected = [item for item in items if item["Status"] == "REJECTED_PRESERVED"]
    assert len(rejected) == 6 and all(item["Reason"] == "UNBALANCED_BRACES" for item in rejected), rejected
    tables = [item for item in rejected if "正文表格" in item["SourceLocation"]]
    assert len(tables) == 2 and len({item["SourceLocation"] for item in tables}) == 2, tables
    assert all(item["Page"] == 2 for item in tables), tables
    headers = [item for item in rejected if "页眉" in item["SourceLocation"]]
    assert len(headers) == 2 and all(item["Page"] is None and "共用" in item["SourceLocation"] for item in headers), headers
    assert any("第 1 节" in item["SourceLocation"] for item in headers)
    assert any("第 2 节" in item["SourceLocation"] for item in headers)
    summary = ui["Summary"]
    assert "原文第 2 页" in summary and "原文第 3 页" in summary and "花括号" in summary, summary
    assert "UNBALANCED_BRACES" not in summary and "UI_ITEM_" not in summary, summary
    assert "Word 位置" not in summary and str(source) not in summary, summary
    if reflow:
        pages = json.loads((folder / "page-probe.json").read_text(encoding="utf-8-sig"))
        assert pages["SourcePage"] == 2 and pages["OutputPage"] > pages["SourcePage"], pages
        assert pages["SourceRestartPage"] == 3 and pages["SourceRestartAdjustedPage"] == 1, pages
        assert pages["OriginalsUnchanged"] is True, pages
        assert failure["Page"] != pages["OutputPage"], (failure, pages)
    # 原式保留由 DOCX 实际文本独立核对，不把转换成功解释成数学等价。
    output = Document(ui["Output"])
    assert "BODY_TWO $x_{1$" in "\n".join(p.text for p in output.paragraphs)
    assert "DIAGNOSTIC_RESTART $x_{1$" in "\n".join(p.text for p in output.paragraphs)
    if reflow:
        assert "FAILURE_TARGET $f+7$" in "\n".join(p.text for p in output.paragraphs)
    assert all("$x_{1$" in table.cell(0, 0).text for table in output.tables)
    assert all("$x_{1$" in section.header.paragraphs[0].text for section in output.sections)


if __name__ == "__main__":
    action, directory = sys.argv[1], Path(sys.argv[2])
    if action == "fixture":
        fixture(directory)
    elif action == "check":
        check(directory, sys.argv[3] == "reflow")
    else:
        raise ValueError(action)
    print("DIAGNOSTICS_HELPER_PASSED: " + action)
