"""原生显示布局与全篇补编号的独立 DOCX 检查；不启动 Office。"""
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Pt
from lxml import etree
from audit_docx import audit

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
NS = {"w": W[1:-1]}


def fixture(folder):
    doc = Document()
    doc.core_properties.author = "Synthetic test fixture"
    doc.core_properties.last_modified_by = ""
    doc.styles["Normal"].font.size = Pt(12)
    style = doc.styles.add_style("SourceSpacing", 1)
    style.base_style = doc.styles["Normal"]
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(9)
    style.paragraph_format.line_spacing = 1.5
    doc.add_paragraph("前文保持。")
    for i, prefix in enumerate(("", "", "", "\t", "\t\t"), 1):
        p = doc.add_paragraph(prefix + r"\[\frac{x_{" + str(i) + r"}+1}{y+2}\]")
        if i == 3:
            p.style = style
        else:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.first_line_indent = Pt(24)
    doc.add_paragraph("行内 $a+b$ 保持同行。")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = r"表格 \[t+1\] 保持。"
    doc.add_paragraph("尾部正文保持。")
    doc.save(folder / "input.docx")


def mixed(source, target):
    with ZipFile(source) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    first = root.xpath("//w:body/w:p[.//*[local-name()='OLEObject']]", namespaces=NS)[0]
    p = etree.Element(W + "p")
    etree.SubElement(etree.SubElement(p, W + "r"), W + "t").text = r"\[z+17\]"
    first.addnext(p)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(target, "w", ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def reference_fixture(source, target):
    with ZipFile(source) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    p = root.xpath("//w:body/w:p[.//*[local-name()='OLEObject']]", namespaces=NS)[0]
    code = p.xpath('.//w:instrText[contains(., "MACROBUTTON MTPlaceRef")]', namespaces=NS)[0]
    run = code.getparent()
    while run.getparent() is not p:
        run = run.getparent()
    index = p.index(run)
    while not p[index].xpath('.//w:fldChar[@w:fldCharType="begin"]', namespaces=NS):
        index -= 1
    begin, depth, end = index, 0, None
    for index in range(begin, len(p)):
        for value in p[index].xpath('.//w:fldChar/@w:fldCharType', namespaces=NS):
            depth += (value == "begin") - (value == "end")
        if depth == 0:
            end = index
            break
    assert end is not None
    mark = etree.Element(W + "bookmarkStart", {W + "id": "900001", W + "name": "NativeDisplayNumber"})
    p.insert(end + 1, etree.Element(W + "bookmarkEnd", {W + "id": "900001"}))
    p.insert(begin, mark)
    new = etree.Element(W + "p")
    etree.SubElement(etree.SubElement(new, W + "r"), W + "t").text = r"\[referenceNew+1\]"
    p.addprevious(new)
    ref = etree.Element(W + "p")
    etree.SubElement(etree.SubElement(ref, W + "r"), W + "t").text = "引用："
    for kind, text in (("begin", None), (None, " REF NativeDisplayNumber \\h "), ("separate", None), (None, "(1)"), ("end", None)):
        r = etree.SubElement(ref, W + "r")
        if kind:
            etree.SubElement(r, W + "fldChar", {W + "fldCharType": kind})
        else:
            etree.SubElement(r, W + ("instrText" if text.startswith(" REF") else "t")).text = text
    body = root.find(W + "body")
    body.insert(len(body) - 1, ref)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(target, "w", ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def legacy_fixture(source, target):
    with ZipFile(source) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    # 文档原有的共享样式可能也被表格使用；只归一化正文目标，不能全局改坏范围外段落。
    styles = etree.fromstring(parts["word/styles.xml"])
    native = styles.xpath('//w:style[w:name/@w:val="MTDisplayEquation"]', namespaces=NS)[0]
    props = native.find(W + "pPr")
    if props is None:
        props = etree.SubElement(native, W + "pPr")
    for name in ("jc", "tabs"):
        for child in props.findall(W + name):
            props.remove(child)
    etree.SubElement(props, W + "jc", {W + "val": "right"})
    stops = etree.SubElement(props, W + "tabs")
    for kind, position in (("center", "1600"), ("right", "3200")):
        etree.SubElement(stops, W + "tab", {W + "val": kind, W + "pos": position})
    table_p = root.find(".//" + W + "tbl/" + W + "tr/" + W + "tc/" + W + "p")
    table_props = table_p.find(W + "pPr")
    if table_props is None:
        table_props = etree.Element(W + "pPr")
        table_p.insert(0, table_props)
    etree.SubElement(table_props, W + "pStyle", {W + "val": "MTDisplayEquation"})
    for p in root.xpath("//w:body/w:p[.//*[local-name()='OLEObject']]", namespaces=NS):
        style = p.find(W + "pPr/" + W + "pStyle")
        if style is None or style.get(W + "val") != "MTDisplayEquation":
            continue
        style.set(W + "val", "Normal")
        for tab in p.xpath(".//w:tab", namespaces=NS):
            tab.getparent().remove(tab)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts["word/styles.xml"] = etree.tostring(styles, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(target, "w", ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def effective(paragraph, name):
    value = getattr(paragraph.paragraph_format, name)
    style = paragraph.style
    while value is None and style is not None:
        value = getattr(style.paragraph_format, name)
        style = style.base_style
    return value if value is not None or name == "line_spacing" else 0


def preserved_numbers(source, target):
    # 比较完整编号指令及缓存文字，避免默认模式偷偷更新已有序号。
    def numbers(path):
        return [p._p.xpath(".//w:instrText/text() | .//w:fldSimple/@w:instr | .//w:t/text()")
                for p in Document(path).paragraphs
                if p._p.xpath('.//w:instrText[contains(., "MACROBUTTON MTPlaceRef")]')]
    before, after = numbers(source), numbers(target)
    assert before and before == after, (before, after)
    print("EXISTING_NATIVE_NUMBERS_PRESERVED")


def check(source, target, count, start):
    result = audit(target)
    before, after = Document(source), Document(target)
    displays = [p for p in after.paragraphs if len(p._p.xpath(".//*[local-name()='OLEObject']")) == 1 and p.style.name == "MTDisplayEquation"]
    assert len(displays) == count
    numbers = []
    for p in displays:
        sequence = p._p.xpath(".//w:r/w:tab | .//*[local-name()='OLEObject']")
        assert len(sequence) >= 2 and sequence[0].tag == W + "tab" and sequence[1].tag.endswith("}OLEObject"), p.text
        tabs = list(p.paragraph_format.tab_stops)
        style = p.style
        while not tabs and style is not None:
            tabs = list(style.paragraph_format.tab_stops)
            style = style.base_style
        assert any(t.alignment == 1 for t in tabs) and any(t.alignment == 2 for t in tabs), p._p.xml
        assert effective(p, "first_line_indent") == 0
        assert (effective(p, "alignment") or 0) == 0
        instructions = p._p.xpath(".//w:instrText/text()")
        has_number = any("MACROBUTTON MTPlaceRef" in value for value in instructions)
        assert has_number == (start > 0), instructions
        if has_number:
            # 外层 MACROBUTTON 的文字和嵌套序列缓存共同构成可见编号。
            values = p._p.xpath(".//w:t/text() | .//w:instrText/text()")
            numbers.extend(value.strip() for value in values if value.strip().isdigit())
    assert numbers == list(map(str, range(start, start + count))) if start else not numbers
    # 源段落的行距和段距继续有效，包括通过自定义样式继承的值。
    originals = [p for p in before.paragraphs if p.text.lstrip().startswith(r"\[")]
    if len(originals) == count:
        for old, new in zip(originals, displays):
            for name in ("space_before", "space_after", "line_spacing", "left_indent", "right_indent"):
                assert effective(old, name) == effective(new, name), (name, effective(old, name), effective(new, name))
    assert result["Tables"] == 1 and r"表格 \[t+1\] 保持。" in result["Text"]
    old_table, new_table = before.tables[0].cell(0, 0).paragraphs[0], after.tables[0].cell(0, 0).paragraphs[0]
    assert effective(old_table, "alignment") == effective(new_table, "alignment"), "范围外表格的共享样式被改变"
    print(json.dumps({"Status": "PASSED", "NativeDisplayParagraphs": count, "VisibleNumbers": numbers, "TablesPreserved": True}, ensure_ascii=False))


if __name__ == "__main__":
    mode = sys.argv[1]
    if mode == "fixture":
        fixture(Path(sys.argv[2]))
    elif mode == "mixed":
        mixed(Path(sys.argv[2]), Path(sys.argv[3]))
    elif mode == "reference":
        reference_fixture(Path(sys.argv[2]), Path(sys.argv[3]))
    elif mode == "legacy":
        legacy_fixture(Path(sys.argv[2]), Path(sys.argv[3]))
    elif mode == "preserved-numbers":
        preserved_numbers(Path(sys.argv[2]), Path(sys.argv[3]))
    elif mode == "simple-layout":
        before, after = map(Document, sys.argv[2:4])
        assert len(before.paragraphs) == len(after.paragraphs)
        for old, new in zip(before.paragraphs, after.paragraphs):
            for name in ("space_before", "space_after", "line_spacing"):
                assert effective(old, name) == effective(new, name), name
            assert new.style.name == "MTDisplayEquation" and (effective(new, "alignment") or 0) == 0
            assert effective(new, "first_line_indent") == 0
        print("NATIVE_DISPLAY_SIMPLE_LAYOUT_PASSED")
    elif mode == "check-reference":
        doc = Document(sys.argv[2])
        p = [p for p in doc.paragraphs if p.text.startswith("引用：")][0]
        assert p.text == "引用：(2)", p.text
        assert p._p.xpath('.//w:instrText[contains(., "REF NativeDisplayNumber")]')
        print("NATIVE_NUMBER_REFERENCE_UPDATED")
    elif mode == "manual":
        import fitz
        path = Path(sys.argv[2])
        document = Document(path)
        with fitz.open(str(path) + ".pdf") as pdf:
            hits = [rect for page in pdf for rect in page.search_for("(1)")]
        right = document.sections[0].page_width.pt - document.sections[0].right_margin.pt
        assert len(hits) == 1 and abs(hits[0].x1 - right) <= 1, (hits, right)
        print(json.dumps({"Status": "PASSED", "NativeInsertNumberRightAligned": True}))
    else:
        check(Path(sys.argv[2]), Path(sys.argv[3]), int(sys.argv[4]), int(sys.argv[5]))
