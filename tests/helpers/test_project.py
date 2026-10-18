"""检查公开 Python 语法，并验证合成样本结构与匿名元数据。"""
import ast
import runpy
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from docx import Document

root = Path(__file__).resolve().parents[2]
for path in (root / 'tests/helpers').glob('*.py'):
    ast.parse(path.read_text(encoding='utf-8-sig'), filename=path.name)
with TemporaryDirectory(prefix='fixture-check-', dir=root / 'tests/output') as folder:
    sys.argv = ['test_lightweight_fixtures.py', folder]
    runpy.run_path(str(root / 'tests/helpers/test_lightweight_fixtures.py'), run_name='__main__')
    mixed = Document(Path(folder) / 'mixed.docx')
    assert len(mixed.tables) == 1 and len(mixed.paragraphs) == 8
    assert '$x_1+1$' in mixed.paragraphs[0].text
    assert mixed.core_properties.author == 'Synthetic test fixture'
    assert mixed.core_properties.last_modified_by == ''
    with ZipFile(Path(folder) / 'mixed.docx') as package:
        assert not any(name.startswith('word/embeddings/') for name in package.namelist())
print('PYTHON_AND_SYNTHETIC_FIXTURES_PASSED')
