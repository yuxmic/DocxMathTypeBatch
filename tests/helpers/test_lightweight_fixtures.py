"""生成轻量化回归文档；只写调用方指定的测试目录。"""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

root = Path(sys.argv[1]).resolve()
root.mkdir(parents=True, exist_ok=True)
document = Document()
document.core_properties.author = "Synthetic test fixture"
document.core_properties.last_modified_by = ""
document.styles['Normal'].font.size = Pt(12)
document.add_paragraph('保留正文；行内 $x_1+1$ 结束。')
p = document.add_paragraph(r'\[\frac{a+1}{b}\]')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Pt(24)
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(9)
p.runs[0].font.size = Pt(16)
# 段落符和软换行都放在 TeX 分行符之后，右定界符必须完整取得。
document.add_paragraph(r'\[f(x)=\begin{cases}x^2,&x>0\\')
document.add_paragraph(r'-x,&x\leq0\end{cases}\]')
document.add_paragraph('$$' + r'\begin{array}{cc}a&b\\' + '\n' + r'c&d\end{array}' + '$$')
document.add_paragraph('$' + r'\begin{cases}1,&x>0\\')
document.add_paragraph(r'0,&x\leq0\end{cases}' + '$')
table = document.add_table(rows=1, cols=1)
table.cell(0, 0).text = r'表格 $t_1$ 与 \[t_2\] 保留。'
document.add_paragraph('结尾正文保持。')
document.save(root / 'mixed.docx')

# 破损环境、货币和后续正常公式不得互相吞并。
document = Document()
document.core_properties.author = "Synthetic test fixture"
document.core_properties.last_modified_by = ""
for text in ('Budget $5', 'paid $10', r'\[\begin{cases}a\\', r'b\end{array}\]', r'\[missing', r'\[q+2\]', r'$z_1$'):
    document.add_paragraph(text)
document.save(root / 'boundaries.docx')

# 简单布局样本同时用于恢复注入，保留三条公式的数量约定。
document = Document()
document.core_properties.author = "Synthetic test fixture"
document.core_properties.last_modified_by = ""
for alignment in (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
    p = document.add_paragraph(r'\[x+1\]')
    p.alignment = alignment
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9)
document.save(root / 'layout.docx')

# 居中回归覆盖字符/磅值缩进及前导 Tab，正文、表格和段距必须保留。
document = Document()
document.core_properties.author = "Synthetic test fixture"
document.core_properties.last_modified_by = ""
document.styles['Normal'].font.size = Pt(12)
for name, alignment, indent, tab, left, right in (
    ('left', 0, 0, False, 0, 0), ('justify-indent', 3, 24, False, 0, 0),
    ('right-indent', 2, 24, False, 0, 0), ('center-indent', 1, 24, False, 0, 0),
    ('leading-tab', 0, 0, True, 0, 0), ('hanging', 0, -18, False, 18, 0),
    ('side-indents', 3, 24, False, 18, 12), ('character-indent', 3, 24, False, 0, 0),
):
    document.add_paragraph('before-' + name + '; inline $a_1+1$ remains inline.')
    p = document.add_paragraph(('\t' if tab else '') + r'\[\frac{x_1+1}{y+2}\]')
    p.alignment = alignment
    f = p.paragraph_format
    f.first_line_indent, f.left_indent, f.right_indent = Pt(indent), Pt(left), Pt(right)
    f.space_before, f.space_after, f.line_spacing = Pt(6), Pt(9), 1.25
    p.runs[0].font.size = Pt(14 if name == 'right-indent' else 12)
    if name == 'character-indent':
        p._p.pPr.ind.set(qn('w:firstLineChars'), '200')
    document.add_paragraph('after-' + name + ' unchanged.')
document.add_table(rows=1, cols=1).cell(0, 0).text = r'untouched $t_1$ and \[t_2\]'
document.add_paragraph('End remains unchanged.')
document.save(root / 'alignment.docx')

# 定界符与段落上下文独立判断；普通括号、转义及破损标记不得误识别。
document = Document()
document.core_properties.author = "Synthetic test fixture"
document.core_properties.last_modified_by = ""
for text in (r'before $$a+1$$ after', r'before \[b+1\] after', r'before \(c+1\) after',
             r'$d+1$', r'$$e+1$$', r'\[f+1\]', r'\(g+1\)', r'[h+1\]',
             r'$a$ and $a$', r'\\(not_math\\)', r'\[\sqrt[3]{x}\]',
             r'正文 [ordinary] \] 保留', r'[ordinary]', r'\(broken\]', r'\(y+1\)'):
    document.add_paragraph(text)
document.save(root / 'delimiters.docx')

# aligned 兼容必须保留多组等式、嵌套矩阵及 Word 源码换行。
document = Document()
document.core_properties.author = "Synthetic test fixture"
document.core_properties.last_modified_by = ""
document.styles['Normal'].font.size = Pt(12)
document.add_paragraph(r'\[\begin{aligned}a_1&=b_1\\a_2&=b_2\end{aligned}\]')
document.add_paragraph(r'\[\begin{aligned}P_i&=V_i\sum_{j=1}^{n}V_j(G_{ij}\cos\theta_{ij}+B_{ij}\sin\theta_{ij}),&Q_i&=V_i\sum_{j=1}^{n}V_j(G_{ij}\sin\theta_{ij}-B_{ij}\cos\theta_{ij})\\')
document.add_paragraph(r'P_2&=Q_2,&P_3&=Q_3\end{aligned}\]')
document.add_paragraph(r'\[\begin{aligned}[t]c_1&=d_1\\c_2&=d_2\end{aligned}\]')
document.add_paragraph(r'\[\begin{aligned}e_1&=\begin{matrix}m_1&m_2\\m_3&m_4\end{matrix}\end{aligned}\]')
document.add_paragraph(r'\[\begin{aligned}f_1&=\left(\begin{aligned}n_1&=n_2&n_3&=n_4\\n_5&=n_6&n_7&=n_8\end{aligned}\right)\end{aligned}\]')
document.add_paragraph(r'\[\begin{array}{rl}q_1&=q_2\\q_3&=q_4\end{array}\]')
document.add_paragraph(r'inline \(\begin{aligned}u_1&=u_2\\u_3&=u_4\end{aligned}\) remains inline.')
document.save(root / 'aligned.docx')

# 性能样本含重复公式和普通段落，旧、新成品使用同一输入。
for count, paragraphs in ((20, 80), (100, 400), (10, 2000)):
    document = Document()
    document.core_properties.author = "Synthetic test fixture"
    document.core_properties.last_modified_by = ""
    document.styles['Normal'].font.size = Pt(12)
    for index in range(paragraphs):
        document.add_paragraph(r'\[x_1+1\]' if index % (paragraphs // count) == 0 else f'普通正文 {index} ABC 123。')
    document.save(root / f'performance-{count}-{paragraphs}.docx')
print(root)
