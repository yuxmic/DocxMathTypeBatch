"""核对不对称定界符、百分号和非 OLE 嵌入包，并保留同文档内的失败原式。"""
import json
import sys
from collections import Counter
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.shared import Pt
from audit_docx import audit

COMMENT = r'\[\begin{aligned}p&=1% ignored' + '\n' + r'\\q&=2\end{aligned}\]'


def fixture(folder):
    document = Document()
    document.core_properties.author = 'Synthetic test fixture'
    document.core_properties.last_modified_by = ''
    for value in (r'[xy\]', r' [x\] ', r'[z_{1}\]', '$i_1$', '$$d_1$$', r'\(p_1\)', r'\[q_1\]', r'[\]', 'End stays exact.'):
        document.add_paragraph(value)
    document.save(folder / 'reverse-input-forms.docx')

    # 真正挂接图表与工作簿关系，避免 Word 清理没有引用的合成嵌入部件。
    spreadsheet = BytesIO()
    with ZipFile(spreadsheet, 'w', ZIP_DEFLATED) as package:
        package.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        package.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        package.writestr('xl/workbook.xml', '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>')
        package.writestr('xl/_rels/workbook.xml.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        package.writestr('xl/worksheets/sheet1.xml', '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Value</t></is></c><c r="B1"><v>25</v></c></row></sheetData></worksheet>')
    document = Document()
    document.core_properties.author = 'Synthetic test fixture'
    document.core_properties.last_modified_by = ''
    document.styles['Normal'].font.size = Pt(12)
    document.add_paragraph(r'\[50\%\]')
    document.add_paragraph(r'Keep unsupported $\text{汉字}$ exact.')
    document.add_paragraph(r'\[unfinished')
    document.add_table(rows=1, cols=1).cell(0, 0).text = r'Cell $25\%$ and $t_1$ stay here.'
    # 注释既不能排成可见公式，也不能在换行合并时丢失原文。
    document.add_paragraph(COMMENT)
    document.tables[0].add_row().cells[0].text = COMMENT
    workbook = Part(PackURI('/word/embeddings/input-forms.xlsx'), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', spreadsheet.getvalue(), document.part.package)
    chart = Part(PackURI('/word/charts/chart1.xml'), 'application/vnd.openxmlformats-officedocument.drawingml.chart+xml', b'''<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <c:chart><c:plotArea><c:pieChart><c:varyColors val="0"/><c:ser><c:idx val="0"/><c:order val="0"/><c:cat><c:strRef><c:f>Sheet1!$A$1</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>Value</c:v></c:pt></c:strCache></c:strRef></c:cat><c:val><c:numRef><c:f>Sheet1!$B$1</c:f><c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="1"/><c:pt idx="0"><c:v>25</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser></c:pieChart></c:plotArea><c:plotVisOnly val="1"/></c:chart>
      <c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData></c:chartSpace>''', document.part.package)
    assert chart.relate_to(workbook, RT.PACKAGE) == 'rId1'
    relation = document.part.relate_to(chart, RT.CHART)
    document.add_paragraph().add_run()._r.append(parse_xml(f'''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><wp:inline><wp:extent cx="2743200" cy="1828800"/><wp:docPr id="1" name="Synthetic chart"/><wp:cNvGraphicFramePr/><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart r:id="{relation}"/></a:graphicData></a:graphic></wp:inline></w:drawing>'''))
    document.add_paragraph('End stays exact.')
    document.save(folder / 'embedded-input-forms.docx')


def check(folder):
    for delimiter in ('bracket', 'dollar'):
        path = folder / ('reverse-forms-' + delimiter + '.docx')
        paragraphs = [paragraph.text for paragraph in Document(path).paragraphs]
        opening, closing = (r'\[', r'\]') if delimiter == 'bracket' else ('$$', '$$')
        assert paragraphs == [opening+'xy'+closing, ' '+opening+'x'+closing+' ', opening+'z_{1}'+closing, '$i_1$', opening+'d_1'+closing, '$p_1$', opening+'q_1'+closing, r'[\]', 'End stays exact.'], paragraphs
        report = json.loads(path.with_suffix('.json').read_text(encoding='utf-8-sig'))
        assert report['Succeeded'] == 7 and report['RejectedPreserved'] == 1 and report['Status'] == 'PARTIAL'
    source = folder / 'embedded-input-forms.docx'
    with ZipFile(source) as package:
        embeddings = Counter(package.read(name) for name in package.namelist() if name.startswith('word/embeddings/'))
    for name in ('embedded-forms', 'embedded-forms-extra', 'embedded-forms-reverse'):
        path = folder / (name + '.docx')
        report = json.loads(path.with_suffix('.json').read_text(encoding='utf-8-sig'))
        document = Document(path)
        assert report['Status'] == 'PARTIAL' and report['WordReopenVerified'] and report['OriginalFileBytesUnchanged']
        assert document.paragraphs[1].text == r'Keep unsupported $\text{汉字}$ exact.'
        assert document.paragraphs[2].text == r'\[unfinished' and document.paragraphs[-1].text == 'End stays exact.'
        assert len(document.element.xpath('//*[local-name()="chart"]')) == 1
        with ZipFile(path) as package:
            assert Counter(package.read(part) for part in package.namelist() if part.startswith('word/embeddings/') and part.endswith('.xlsx')) == embeddings
        result = audit(path)
        if name == 'embedded-forms-reverse':
            assert result['MathType'] == 0 and report['ExportedMathType'] == 1
            assert r'50\%' in document.paragraphs[0].text and r'25\%' in document.tables[0].cell(0, 0).text
        else:
            assert report['FailedPreserved'] == 1 and report['RejectedPreserved'] == (3 if name == 'embedded-forms-extra' else 2)
            comments = [item for item in report['Items'] if item['Reason'] == 'UNSUPPORTED_TEX_COMMENT']
            assert len(comments) == (2 if name == 'embedded-forms-extra' else 1)
            assert report['Succeeded'] == result['MathType'] == (1 if name == 'embedded-forms' else 3)
            if name == 'embedded-forms':
                assert document.tables[0].cell(0, 0).text == r'Cell $25\%$ and $t_1$ stay here.'
        assert COMMENT in [paragraph.text for paragraph in document.paragraphs]
        assert document.tables[0].cell(1, 0).text == COMMENT
    print('INPUT_FORMS_AND_RETAINED_EMBEDDINGS_PASSED')


def shared_fixture(folder):
    # 两个对象引用同一 OLE 与预览部件，分别覆盖全部导出和表格中仍保留一个引用。
    for table in (False, True):
        document = Document(folder / 'embedded-forms.docx')
        paragraph = deepcopy(document.paragraphs[0]._p)
        for node in paragraph.iter():
            for attribute in ('id', 'ShapeID', 'ObjectID'):
                if attribute in node.attrib:
                    node.set(attribute, '_shared_fixture')
        if table:
            cell = document.tables[0].add_row().cells[0]
            cell._tc.replace(cell.paragraphs[0]._p, paragraph)
        else:
            document.element.body.insert(1, paragraph)
        document.save(folder / ('shared-table-input.docx' if table else 'shared-input.docx'))


def shared_check(folder):
    for table in (False, True):
        name = 'shared-table' if table else 'shared'
        source = folder / (name + '-input.docx')
        target = folder / (name + '.docx')
        report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
        assert report['Status'] == 'PARTIAL' and report['ExportedMathType'] == (1 if table else 2)
        assert report['WordReopenVerified'] and report['OriginalFileBytesUnchanged']
        assert audit(target)['MathType'] == int(table)
        if table:
            with ZipFile(source) as original, ZipFile(target) as saved:
                before = [original.read(part) for part in original.namelist() if part.startswith('word/embeddings/') and part.endswith('.bin')]
                after = [saved.read(part) for part in saved.namelist() if part.startswith('word/embeddings/') and part.endswith('.bin')]
                assert before and after and all(value in before for value in after)
    print('SHARED_REVERSE_ASSETS_PASSED')


def boundary_fixture(folder):
    # 行外空白可属于控件或空结果域；LaTeX 和 OMML 都不得把它们纳入替换。
    document = Document()
    document.core_properties.author = 'Synthetic test fixture'
    document.core_properties.last_modified_by = ''
    word = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    math = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    index = 0
    for omml in (False, True):
        for field in (False, True):
            for before in (True, False):
                index += 1
                paragraph = document.add_paragraph()
                boundary = parse_xml(f'<w:fldSimple xmlns:w="{word}" w:instr="SEQ Hidden \\h"/>' if field else
                                     f'<w:sdt xmlns:w="{word}"><w:sdtPr><w:id w:val="{index}"/><w:text/></w:sdtPr><w:sdtContent><w:r><w:t xml:space="preserve"> </w:t></w:r></w:sdtContent></w:sdt>')
                if before:
                    paragraph._p.append(boundary)
                if omml:
                    paragraph._p.append(parse_xml(f'<m:oMathPara xmlns:m="{math}"><m:oMath><m:r><m:t>x+{index}</m:t></m:r></m:oMath></m:oMathPara>'))
                else:
                    paragraph.add_run(r'\[x+' + str(index) + r'\]')
                if not before:
                    paragraph._p.append(boundary)
    paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph._p.append(parse_xml(f'<w:sdt xmlns:w="{word}"><w:sdtPr><w:id w:val="9"/><w:text/></w:sdtPr><w:sdtContent><w:r><w:t xml:space="preserve"> </w:t></w:r></w:sdtContent></w:sdt>'))
    paragraph.add_run(r'\[z+1\]')
    # 普通前导空白和完整跨段公式仍沿用原有转换规则。
    document.add_paragraph('\t ' + r'\[u+v\]' + ' \t')
    document.add_paragraph(r'\[\begin{aligned}a&=b\\')
    document.add_paragraph(r'c&=d\end{aligned}\]')
    document.add_paragraph('End stays exact.')
    document.save(folder / 'formula-boundaries-input.docx')


def boundary_check(folder):
    target = folder / 'formula-boundaries.docx'
    report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
    assert report['Status'] == 'COMPLETE' and report['Succeeded'] == 11
    assert report['Experimental']['Succeeded'] == 1 and not report['Issues']
    assert report['OriginalFileBytesUnchanged'] and report['WordReopenVerified']
    document = Document(target)
    word = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    controls = document.element.xpath('//w:sdt')
    assert len(controls) == 5 and all(''.join(text.text or '' for text in node.iter(word + 't')) == ' ' for node in controls)
    instructions = [node.text or '' for node in document.element.xpath('//w:instrText')]
    instructions += [node.get(word + 'instr') for node in document.element.xpath('//w:fldSimple')]
    assert len([code for code in instructions if 'SEQ Hidden' in code and r'\h' in code]) == 4
    assert document.paragraphs[-1].text == 'End stays exact.'
    assert audit(target)['MathType'] == 11
    print('FORMULA_BOUNDARIES_PASSED')


def fast_check(folder):
    # 外部测试可以审计样本；产品快速模式本身必须没有重开、重扫或资产核验。
    for name, count in (('formula-boundaries-fast', 11), ('formula-boundaries-fast-reverse', 11), ('formula-boundaries-fast-failure', 10)):
        target = folder / (name + '.docx')
        report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
        assert report['FastConversion'] and report['Validation'] == 'SKIPPED_EXPERIMENTAL_FAST'
        assert report['PostConversionScans'] == 0 and not report['WordReopenVerified']
        assert report['OriginalFileBytesUnchanged'] is None and report['ExistingNativeStreamsUnchanged'] is None
        assert report['Succeeded'] == count
        assert audit(target)['MathType'] == (0 if name.endswith('reverse') else count)
        log = target.with_suffix('.log').read_text(encoding='utf-8')
        assert not any(marker in log for marker in ('UI_PHASE VERIFYING', 'UI_PHASE REOPENING', 'UI_PHASE RESCANNING', 'UI_PHASE PACKAGE_VERIFY', 'UI_PHASE PACKAGE_COMPARE', 'UI_STAGE_PROGRESS INPUT_PARTS', 'UI_STAGE_PROGRESS SCAN_RANGES', 'UI_EXPERIMENTAL_PHASE VERIFYING', 'EXPERIMENTAL_CONVERSION_VERIFIED', 'TIMING_STAGE BIND_OBJECTS', 'TIMING_STAGE SOURCE_SNAPSHOT', 'TIMING_STAGE SOURCE_COMPARE', 'BODY_CHECKPOINT_RESTORED'))
        if name.endswith('failure'):
            assert report['Status'] == 'PARTIAL' and report['FailedUnverified'] == 1
            assert any(item['Status'] == 'FAILED_UNVERIFIED' for item in report['Items'])
        else:
            assert report['Status'] == 'COMPLETE' and report['FailedUnverified'] == 0
    print('FAST_CONVERSION_UNVERIFIED_OUTPUT_PASSED')


def manual_omml_fixture(folder):
    """合成行内原生公式加手动编号；同时保留正文、单位、控件和普通域反例。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from test_experimental import math_run

    document = Document()
    document.core_properties.author = 'Synthetic test fixture'
    document.core_properties.last_modified_by = ''
    for index, number in enumerate(('  （27）', '\t(3-2)  '), 1):
        paragraph = document.add_paragraph('  ')
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph._p.append(math_run(index))
        paragraph.add_run(number)
        if index == 1:
            paragraph.add_run().add_break()
    paragraph = document.add_paragraph('Inline text ')
    paragraph._p.append(math_run(3))
    paragraph.add_run(' (99) remains.')
    paragraph = document.add_paragraph()
    paragraph._p.append(math_run(4))
    paragraph.add_run(' (m/s)')
    paragraph = document.add_paragraph()
    paragraph._p.append(math_run(5))
    paragraph._p.append(parse_xml('<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:sdtPr/><w:sdtContent><w:r><w:t xml:space="preserve"> </w:t></w:r></w:sdtContent></w:sdt>'))
    paragraph.add_run(' (88)')
    paragraph = document.add_paragraph()
    paragraph._p.append(math_run(6))
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), r'SEQ ManualBoundary \h')
    paragraph._p.append(field)
    paragraph.add_run(' (89)')
    document.save(folder / 'manual-omml-input.docx')


def manual_latex_fixture(folder):
    """复用相同段落边界，替换为四种 LaTeX 定界符，覆盖数字编号及非编号反例。"""
    from docx.oxml import OxmlElement

    manual_omml_fixture(folder)
    document = Document(folder / 'manual-omml-input.docx')
    formulas = (r'\[\frac{x+1}{2}\]', '$$x+2$$', '$x+3$', r'\(x+4\)', '$x+5$', r'\(x+6\)')
    for paragraph, formula in zip(document.paragraphs, formulas):
        math = paragraph._p.xpath('.//m:oMath')[0]
        run = OxmlElement('w:r')
        node = OxmlElement('w:t')
        node.text = formula
        run.append(node)
        math.addprevious(run)
        math.getparent().remove(math)
    document.add_paragraph('  $x+7$  (3.2)')
    document.add_paragraph(r'  \(x+8\)  （4－1）')
    document.save(folder / 'manual-latex-input.docx')


def manual_equation_check(folder, kind):
    """独立检查旧编号清除、原生居中制表位，以及应保留的非编号文字和边界。"""
    from docx.oxml.ns import qn
    from native_number_text import native_visible_numbers
    from test_native_display import effective

    prefix = f'manual-{kind}'
    positive = [0, 1, 6, 7] if kind == 'latex' else [0, 1]
    expected = 8 if kind == 'latex' else 6
    for name, numbers in ((prefix, 0), (prefix + '-fast', 0), (prefix + '-numbered', len(positive))):
        target = folder / (name + '.docx')
        report = json.loads(target.with_suffix('.json').read_text(encoding='utf-8-sig'))
        assert report['Status'] == 'COMPLETE' and report['Succeeded'] == expected, report
        assert report['NumbersAdded'] == numbers, report
        assert [item['Display'] for item in report['Items']] == [i in positive for i in range(expected)]
        document = Document(target)
        paragraphs = document.paragraphs
        for index in positive:
            paragraph = paragraphs[index]
            body = paragraph._p
            text = ''.join(body.xpath('.//w:t/text()'))
            assert not any(old in text for old in ('27', '3-2', '3.2', '4－1')), (name, text)
            if not numbers:
                assert not text.strip(), (name, text)
            assert not body.xpath('.//w:br'), (name, 'soft break remains')
            assert len(body.xpath('./w:r/w:tab')) == 2, (name, 'display tabs')
            tabs = body.xpath('./w:pPr/w:tabs/w:tab') or paragraph.style.element.xpath('./w:pPr/w:tabs/w:tab')
            assert [node.get(qn('w:val')) for node in tabs] == ['center', 'right'], (name, 'centering')
            assert paragraph.style.name == 'MTDisplayEquation' and effective(paragraph, 'first_line_indent') == 0
            assert (effective(paragraph, 'alignment') or 0) == 0
        assert '(99) remains.' in paragraphs[2].text and '(m/s)' in paragraphs[3].text
        assert '(88)' in paragraphs[4].text and paragraphs[4]._p.xpath('.//w:sdt')
        assert '(89)' in paragraphs[5].text
        instructions = document.element.xpath('//w:instrText/text()') + document.element.xpath('//w:fldSimple/@w:instr')
        assert any('SEQ ManualBoundary' in code for code in instructions)
        assert len(native_visible_numbers(document.element)) == numbers
        assert audit(target)['MathType'] == expected
    print(f'MANUAL_{kind.upper()}_NUMBERS_AND_LAYOUT_PASSED')


def manual_omml_recovery_check(folder):
    """失败原式的空白、手动编号和软换行也应保留，其他公式正常输出。"""
    report = json.loads((folder / 'manual-omml-recovery.json').read_text(encoding='utf-8-sig'))
    assert report['Status'] == 'PARTIAL' and report['Succeeded'] == 5, report
    before = Document(folder / 'manual-omml-input.docx').paragraphs[0]
    output = Document(folder / 'manual-omml-recovery.docx')
    after = output.paragraphs[0]
    assert before.text == after.text and before.alignment == after.alignment
    assert len(after._p.xpath('.//m:oMath')) == 1
    assert audit(folder / 'manual-omml-recovery.docx')['MathType'] == 5
    print('MANUAL_OMML_RECOVERY_PASSED')


if __name__ == '__main__':
    mode, directory = sys.argv[1:3]
    {'fixture': fixture, 'check': check, 'shared-fixture': shared_fixture, 'shared-check': shared_check,
     'boundary-fixture': boundary_fixture, 'boundary-check': boundary_check, 'fast-check': fast_check,
     'manual-omml-fixture': manual_omml_fixture, 'manual-omml-check': lambda path: manual_equation_check(path, 'omml'),
     'manual-latex-fixture': manual_latex_fixture, 'manual-latex-check': lambda path: manual_equation_check(path, 'latex'),
     'manual-omml-recovery-check': manual_omml_recovery_check}[mode](Path(directory))
